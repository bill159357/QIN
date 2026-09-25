from __future__ import annotations

import base64
import html
import io
import json
import os
import re
import requests
import shutil
import sqlite3
import threading
import time
import uuid
import importlib
import multiprocessing
import queue
import tempfile
from contextlib import asynccontextmanager
from urllib.parse import quote, unquote, urlsplit
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse, Response
from PIL import Image
from pydantic import BaseModel


def _load_local_env_files() -> None:
    candidates = [
        Path(__file__).resolve().with_name(".env.local"),
        Path(__file__).resolve().with_name(".env"),
    ]
    for env_path in candidates:
        if not env_path.is_file():
            continue
        for raw_line in env_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, raw_value = line.split("=", 1)
            env_key = key.strip()
            if not env_key or env_key in os.environ:
                continue
            env_value = raw_value.strip()
            if len(env_value) >= 2 and env_value[0] == env_value[-1] and env_value[0] in {'"', "'"}:
                env_value = env_value[1:-1]
            os.environ[env_key] = env_value


_load_local_env_files()

from deepjiandu_runtime import normalize_deepjiandu_options, run_deepjiandu_page

PROMPT_DEFAULT = "<image>\n<|grounding|>Convert the document to markdown. "
OCR_MODEL_DEEPSEEK = "DeepSeek-OCR-2"
OCR_MODEL_PADDLE = "PaddleOCR-VL-1.5"
OCR_MODEL_DEEPJIANDU = "DeepJiandu-OCR-v1"
DEEPBOOK_ASSISTANT_NAME = "简衡"
DEEPBOOK_PLATFORM_NAME = "秦简智读——基于大模型的秦简数字化与智能服务平台"
MOONSHOT_BASE_URL = os.getenv("MOONSHOT_BASE_URL", "https://api.moonshot.cn/v1").strip().rstrip("/")
MOONSHOT_CHAT_COMPLETIONS_URL = f"{MOONSHOT_BASE_URL}/chat/completions"
MOONSHOT_CHAT_MODEL = os.getenv("DEEPBOOK_ASSISTANT_MODEL", "kimi-latest").strip() or "kimi-latest"
MOONSHOT_TIMEOUT_SECONDS = 90
ASSISTANT_CONTEXT_DOC_LIMIT = 3
ASSISTANT_CONTEXT_GALLERY_LIMIT = 3
ASSISTANT_MAX_HISTORY_MESSAGES = 12
ASSISTANT_MAX_MESSAGE_CHARS = 4000
ASSISTANT_SYSTEM_PROMPT = f"""
你是“{DEEPBOOK_ASSISTANT_NAME}”，服务于“{DEEPBOOK_PLATFORM_NAME}”。

你的角色与边界：
1. 优先回答秦简、简牍、出土文献、文字释读、纪年制度、文书格式、OCR复核、数字化整理等相关问题。
2. 回答平台使用问题时，要结合本平台模块：简牍藏馆、秦简识别、全文检索、历史任务。
3. 如果系统提供了“平台资料摘录”，必须优先依据这些资料回答，并明确哪些内容来自平台资料，哪些属于一般学术判断。
4. 对残缺字、异体字、释读分歧、时代断定等不确定问题，要说明不确定性，不得编造定论。
5. 回答风格应专业、克制、清楚，默认使用简洁中文；必要时可分点说明。
6. 当用户贴出简文、释文或 OCR 结果时，可帮助做断句、释文说明、疑难字分析、研究线索提示，但不要把推测当作事实。

禁止事项：
- 不要伪造出土编号、篇名、考古结论、学者观点或馆藏来源。
- 不要声称自己已经查看过图片，除非系统提供的平台资料里明确给出相关图像信息。
- 不要脱离秦简主题泛泛而谈；若问题偏离，可简短回答后拉回到平台或简牍研究场景。
""".strip()

PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_ROOT = Path(os.getenv("OCR_INPUT_DIR", PROJECT_ROOT / "ocr_input"))
OUTPUT_ROOT = Path(os.getenv("OCR_OUTPUT_DIR", PROJECT_ROOT / "ocr_output"))
DB_PATH = Path(os.getenv("DEEPBOOK_DB_PATH", PROJECT_ROOT / "deepbook.db"))
JIANDU_GALLERY_ROOT = Path(
    os.getenv(
        "DEEPBOOK_JIANDU_GALLERY_DIR",
        PROJECT_ROOT / "resources" / "deepjiandu" / "dataset" / "deepjiandu_fullpage_eval_test200",
    )
).resolve()
JIANDU_GALLERY_IMAGES_ROOT = JIANDU_GALLERY_ROOT / "images"
JIANDU_GALLERY_ANNOTATIONS_PATH = JIANDU_GALLERY_ROOT / "annotations.json"
JIANDU_GALLERY_MANIFEST_PATH = JIANDU_GALLERY_ROOT / "manifest.json"

INPUT_ROOT.mkdir(parents=True, exist_ok=True)
OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)


@asynccontextmanager
async def lifespan(_: FastAPI):
    _init_db()
    _start_queue_worker()
    try:
        yield
    finally:
        _close_db()


app = FastAPI(title=f"{DEEPBOOK_PLATFORM_NAME} API", version="0.1.0", lifespan=lifespan)

cors_origins = os.getenv(
    "CORS_ORIGINS",
    "http://localhost:5173,http://127.0.0.1:5173",
).split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin for origin in cors_origins if origin],
    allow_methods=["*"],
    allow_headers=["*"],
)


@dataclass
class OCRJob:
    id: str
    status: str = "pending"
    progress: int = 0
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)
    files: list[str] = field(default_factory=list)
    input_paths: list[str] = field(default_factory=list)
    results: list[dict[str, Any]] = field(default_factory=list)
    error: str | None = None
    prompt: str = PROMPT_DEFAULT
    options: dict[str, Any] = field(default_factory=dict)
    pause_requested: bool = False
    page_count: int | None = None


class QueueReorderRequest(BaseModel):
    order: list[str]


class DocumentUpdateRequest(BaseModel):
    edited_markdown: str | None = None
    edited_text: str | None = None


class AssistantChatMessage(BaseModel):
    role: str
    content: str


class AssistantChatRequest(BaseModel):
    messages: list[AssistantChatMessage]
    include_platform_context: bool = True
    temperature: float | None = None


jobs: dict[str, OCRJob] = {}
jobs_lock = threading.Lock()
model_lock = threading.Lock()
db_lock = threading.Lock()
DB_CONN: sqlite3.Connection | None = None
FTS_ENABLED = True
queue_lock = threading.Lock()
queue_event = threading.Event()
job_queue: list[str] = []
running_job_ids: dict[str, str | None] = {
    OCR_MODEL_DEEPSEEK: None,
    OCR_MODEL_PADDLE: None,
    OCR_MODEL_DEEPJIANDU: None,
}
queue_worker_started = False
job_streams: dict[str, list[queue.Queue]] = {}
job_streams_lock = threading.Lock()
MODEL = None
TOKENIZER = None
OCR_HELPERS = None
ocr_helpers_lock = threading.Lock()
jiandu_gallery_lock = threading.Lock()
jiandu_gallery_records: list[dict[str, Any]] | None = None
jiandu_gallery_index: dict[str, dict[str, Any]] | None = None
jiandu_gallery_manifest: dict[str, Any] | None = None


def _init_db() -> None:
    global DB_CONN, FTS_ENABLED
    if DB_CONN is not None:
        return
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL;")
    conn.execute("PRAGMA synchronous=NORMAL;")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id TEXT NOT NULL,
            title TEXT,
            file_name TEXT,
            markdown TEXT,
            text TEXT,
            raw_text TEXT,
            edited_markdown TEXT,
            edited_text TEXT,
            output_dir TEXT,
            output_dir_rel TEXT,
            output_files TEXT,
            boxes_image TEXT,
            cropped_images TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT
        )
        """
    )
    try:
        conn.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS documents_fts
            USING fts5(title, content)
            """
        )
        FTS_ENABLED = True
    except sqlite3.Error:
        FTS_ENABLED = False
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_documents_job_id
        ON documents(job_id)
        """
    )
    _ensure_document_columns(conn)
    conn.commit()
    DB_CONN = conn


def _close_db() -> None:
    global DB_CONN
    conn = DB_CONN
    DB_CONN = None
    if conn is None:
        return
    try:
        conn.close()
    except sqlite3.Error:
        pass


def _ensure_document_columns(conn: sqlite3.Connection) -> None:
    try:
        rows = conn.execute("PRAGMA table_info(documents)").fetchall()
    except sqlite3.Error:
        return
    columns = {row["name"] if isinstance(row, sqlite3.Row) else row[1] for row in rows}
    to_add = {
        "edited_markdown": "TEXT",
        "edited_text": "TEXT",
        "updated_at": "TEXT",
        "ocr_model": "TEXT",
    }
    for column, col_type in to_add.items():
        if column in columns:
            continue
        try:
            conn.execute(f"ALTER TABLE documents ADD COLUMN {column} {col_type}")
        except sqlite3.Error:
            pass


def _db_execute(query: str, params: tuple[Any, ...] = ()) -> sqlite3.Cursor:
    _init_db()
    if DB_CONN is None:
        raise RuntimeError("Database connection not available.")
    with db_lock:
        cursor = DB_CONN.execute(query, params)
        DB_CONN.commit()
    return cursor


def _db_query(query: str, params: tuple[Any, ...] = ()) -> list[sqlite3.Row]:
    _init_db()
    if DB_CONN is None:
        raise RuntimeError("Database connection not available.")
    with db_lock:
        cursor = DB_CONN.execute(query, params)
        rows = cursor.fetchall()
    return rows


def _deserialize_json(value: str | None, fallback: Any) -> Any:
    if not value:
        return fallback
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return fallback


def _parse_iso_timestamp(value: str | None) -> float:
    if not value:
        return time.time()
    try:
        return datetime.fromisoformat(value).timestamp()
    except Exception:
        return time.time()


def _parse_filter_datetime(value: str | None, end_of_day: bool = False) -> datetime | None:
    text = (value or "").strip()
    if not text:
        return None
    normalized = text.replace("Z", "+00:00")
    try:
        if "T" not in normalized and len(normalized) <= 10:
            parsed = datetime.fromisoformat(normalized)
            parsed = parsed.replace(
                hour=23 if end_of_day else 0,
                minute=59 if end_of_day else 0,
                second=59 if end_of_day else 0,
                microsecond=999999 if end_of_day else 0,
            )
        else:
            parsed = datetime.fromisoformat(normalized)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _coerce_bool_option(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "on"}:
        return True
    if text in {"0", "false", "no", "off"}:
        return False
    return default


def _coerce_int_option(value: Any, default: int, min_value: int | None = None) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        parsed = default
    if min_value is not None and parsed < min_value:
        parsed = min_value
    return parsed


def _coerce_float_option(value: Any, default: float, min_value: float | None = None) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        parsed = default
    if min_value is not None and parsed < min_value:
        parsed = min_value
    return parsed


def _coerce_string_option(value: Any, default: str) -> str:
    text = str(value or "").strip()
    return text or default


def _normalize_paddle_markdown_ignore_labels(raw_labels: Any) -> list[str]:
    allowed = {
        "header",
        "header_image",
        "footer",
        "footer_image",
        "number",
        "footnote",
        "aside_text",
    }
    if raw_labels is None:
        return []
    if isinstance(raw_labels, str):
        source = [part.strip() for part in raw_labels.split(",")]
    elif isinstance(raw_labels, (list, tuple, set)):
        source = [str(item).strip() for item in raw_labels]
    else:
        source = [str(raw_labels).strip()]
    normalized: list[str] = []
    for item in source:
        if not item:
            continue
        key = item.lower().replace("-", "_")
        if key in allowed and key not in normalized:
            normalized.append(key)
    return normalized


def _normalize_paddle_layout_shape_mode(value: Any, default: str = "auto") -> str:
    normalized = _coerce_string_option(value, default).strip().lower()
    allowed = {"rect", "quad", "poly", "auto"}
    if normalized in allowed:
        return normalized
    return default


def _normalize_paddle_prompt_label(value: Any, default: str = "ocr") -> str:
    normalized = _coerce_string_option(value, default).strip().lower()
    allowed = {"ocr", "formula", "table", "chart"}
    if normalized in allowed:
        return normalized
    return default


def _normalize_paddle_layout_merge_mode(value: Any, default: str = "large") -> str:
    normalized = _coerce_string_option(value, default).strip().lower()
    allowed = {"large", "small", "union"}
    if normalized in allowed:
        return normalized
    return default


def _resolve_ocr_model(model_name: str | None, strict: bool = False) -> str:
    text = (model_name or "").strip()
    if not text:
        return OCR_MODEL_DEEPSEEK

    normalized = text.lower().replace("_", "-")
    if normalized in {
        "deepseek-ocr-2",
        "deepseek ocr 2",
        "deepseek",
    }:
        return OCR_MODEL_DEEPSEEK
    if normalized in {
        "paddleocr-vl-1.5",
        "paddleocr-vl-1.5b",
        "paddleocr-vl",
        "paddleocr",
        "paddle",
    }:
        return OCR_MODEL_PADDLE
    if normalized in {
        "deepjiandu-ocr-v1",
        "deepjiandu-ocr",
        "deepjiandu-fullpage-ocr",
        "deepjiandu",
        "jiandu",
    }:
        return OCR_MODEL_DEEPJIANDU

    if strict:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported OCR model. "
                "Available options: "
                f"{OCR_MODEL_DEEPSEEK}, {OCR_MODEL_PADDLE}, {OCR_MODEL_DEEPJIANDU}."
            ),
        )
    return OCR_MODEL_DEEPSEEK


def _normalize_pagination(
    page: int | None,
    page_size: int | None,
    *,
    default_page_size: int = 20,
    max_page_size: int = 50,
) -> tuple[int, int]:
    safe_page = max(1, int(page or 1))
    requested_size = page_size if page_size is not None else default_page_size
    safe_page_size = max(1, min(int(requested_size), max_page_size))
    return safe_page, safe_page_size


def _paginate_items(
    items: list[Any], page: int, page_size: int
) -> tuple[list[Any], int, int, bool, bool]:
    total = len(items)
    total_pages = (total + page_size - 1) // page_size if total > 0 else 0
    if total_pages > 0 and page > total_pages:
        page = total_pages
    elif total_pages == 0:
        page = 1
    offset = (page - 1) * page_size
    paged = items[offset : offset + page_size]
    has_prev = page > 1 and total_pages > 0
    has_next = total_pages > 0 and page < total_pages
    return paged, total, total_pages, has_prev, has_next


def _path_is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return True


def _ensure_jiandu_gallery_ready() -> None:
    if not JIANDU_GALLERY_ROOT.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Jiandu gallery dataset directory not found: {JIANDU_GALLERY_ROOT}",
        )
    if not JIANDU_GALLERY_ANNOTATIONS_PATH.is_file():
        raise HTTPException(
            status_code=404,
            detail=f"Jiandu gallery annotations missing: {JIANDU_GALLERY_ANNOTATIONS_PATH}",
        )
    if not JIANDU_GALLERY_IMAGES_ROOT.is_dir():
        raise HTTPException(
            status_code=404,
            detail=f"Jiandu gallery image directory missing: {JIANDU_GALLERY_IMAGES_ROOT}",
        )


def _read_image_size(image_path: Path) -> tuple[int, int]:
    try:
        with Image.open(image_path) as image:
            return int(image.width), int(image.height)
    except Exception:
        return 1, 1


def _normalize_jiandu_boxes(raw_boxes: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_boxes, list):
        return []

    normalized: list[dict[str, Any]] = []
    for raw_box in raw_boxes:
        if not isinstance(raw_box, dict):
            continue
        try:
            x1 = int(raw_box.get("x1", 0))
            y1 = int(raw_box.get("y1", 0))
            x2 = int(raw_box.get("x2", 0))
            y2 = int(raw_box.get("y2", 0))
        except (TypeError, ValueError):
            continue
        normalized.append(
            {
                "index": _coerce_int_option(raw_box.get("index"), len(normalized) + 1, min_value=1),
                "char": str(raw_box.get("char") or "").strip(),
                "x1": x1,
                "y1": y1,
                "x2": x2,
                "y2": y2,
                "w": max(0, _coerce_int_option(raw_box.get("w"), x2 - x1, min_value=0)),
                "h": max(0, _coerce_int_option(raw_box.get("h"), y2 - y1, min_value=0)),
            }
        )
    return normalized


def _load_jiandu_gallery() -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    global jiandu_gallery_records, jiandu_gallery_index, jiandu_gallery_manifest

    if (
        jiandu_gallery_records is not None
        and jiandu_gallery_index is not None
        and jiandu_gallery_manifest is not None
    ):
        return jiandu_gallery_records, jiandu_gallery_index, jiandu_gallery_manifest

    with jiandu_gallery_lock:
        if (
            jiandu_gallery_records is not None
            and jiandu_gallery_index is not None
            and jiandu_gallery_manifest is not None
        ):
            return jiandu_gallery_records, jiandu_gallery_index, jiandu_gallery_manifest

        _ensure_jiandu_gallery_ready()
        try:
            raw_records = json.loads(JIANDU_GALLERY_ANNOTATIONS_PATH.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=500, detail=f"Invalid jiandu gallery annotations: {exc}") from exc

        if not isinstance(raw_records, list):
            raise HTTPException(status_code=500, detail="Jiandu gallery annotations must be a JSON array.")

        manifest: dict[str, Any] = {}
        if JIANDU_GALLERY_MANIFEST_PATH.is_file():
            try:
                manifest_value = json.loads(JIANDU_GALLERY_MANIFEST_PATH.read_text(encoding="utf-8"))
                if isinstance(manifest_value, dict):
                    manifest = manifest_value
            except json.JSONDecodeError:
                manifest = {}

        records: list[dict[str, Any]] = []
        index: dict[str, dict[str, Any]] = {}
        for raw_record in raw_records:
            if not isinstance(raw_record, dict):
                continue
            item_id = str(raw_record.get("id") or "").strip()
            image_rel_path = str(raw_record.get("image_path") or "").replace("\\", "/").lstrip("/")
            if not item_id or not image_rel_path:
                continue
            image_path = (JIANDU_GALLERY_ROOT / image_rel_path).resolve()
            if not _path_is_within(image_path, JIANDU_GALLERY_ROOT):
                continue

            image_width, image_height = _read_image_size(image_path)
            boxes = _normalize_jiandu_boxes(raw_record.get("boxes"))
            reference_text = re.sub(r"\s+", "", str(raw_record.get("reference_text") or "")).strip()
            num_chars = _coerce_int_option(raw_record.get("num_chars"), len(boxes), min_value=0)

            record = {
                "id": item_id,
                "title": f"简牍 {item_id}",
                "split": str(raw_record.get("split") or "test").strip() or "test",
                "image_file": str(raw_record.get("image_file") or image_path.name).strip() or image_path.name,
                "image_path": image_rel_path,
                "source_image": str(raw_record.get("source_image") or "").strip(),
                "source_xml": str(raw_record.get("source_xml") or "").strip(),
                "num_chars": num_chars,
                "reference_text": reference_text,
                "translation_text": "",
                "translation_note": "当前评测样例未附现代汉语译文，先展示数据集提供的参考释文。",
                "boxes": boxes,
                "box_count": len(boxes),
                "image_width": image_width,
                "image_height": image_height,
                "_image_path": image_path,
            }
            records.append(record)
            index[item_id] = record

        jiandu_gallery_records = records
        jiandu_gallery_index = index
        jiandu_gallery_manifest = manifest
    return jiandu_gallery_records, jiandu_gallery_index, jiandu_gallery_manifest


def _get_jiandu_gallery_item(item_id: str) -> dict[str, Any]:
    _, index, _ = _load_jiandu_gallery()
    normalized_id = str(item_id or "").strip()
    if not normalized_id:
        raise HTTPException(status_code=400, detail="Jiandu gallery item id is required.")
    record = index.get(normalized_id)
    if not record:
        raise HTTPException(status_code=404, detail=f"Jiandu gallery item not found: {normalized_id}")
    return record


def _serialize_jiandu_gallery_item(
    record: dict[str, Any], request: Request, *, include_boxes: bool = False
) -> dict[str, Any]:
    public_api_base = _resolve_public_api_base(request)
    item_id = record["id"]
    reference_text = record.get("reference_text") or ""
    box_count = int(record.get("box_count") or len(record.get("boxes") or []))
    payload = {
        "id": item_id,
        "title": record.get("title") or f"简牍 {item_id}",
        "split": record.get("split") or "test",
        "image_file": record.get("image_file") or "",
        "num_chars": int(record.get("num_chars") or box_count),
        "box_count": box_count,
        "reference_text": reference_text,
        "preview_text": reference_text[:48],
        "translation_text": record.get("translation_text") or "",
        "translation_note": record.get("translation_note") or "",
        "image_width": int(record.get("image_width") or 1),
        "image_height": int(record.get("image_height") or 1),
        "source_image": record.get("source_image") or "",
        "source_xml": record.get("source_xml") or "",
        "image_url": f"{public_api_base}/jiandu/gallery/items/{quote(item_id)}/image",
        "thumbnail_url": f"{public_api_base}/jiandu/gallery/items/{quote(item_id)}/thumbnail",
        "detail_url": f"{public_api_base}/jiandu/gallery/items/{quote(item_id)}",
    }
    if include_boxes:
        payload["boxes"] = record.get("boxes") or []
    return payload


def _render_jiandu_gallery_thumbnail(item_id: str, width: int = 480) -> Response:
    record = _get_jiandu_gallery_item(item_id)
    image_path = record["_image_path"]
    target_width = max(120, min(int(width or 480), 960))
    try:
        with Image.open(image_path) as image:
            working = image.convert("RGB")
            if working.width > target_width:
                target_height = max(1, round(working.height * (target_width / working.width)))
                resampling = getattr(getattr(Image, "Resampling", Image), "LANCZOS")
                working = working.resize((target_width, target_height), resampling)
            buffer = io.BytesIO()
            working.save(buffer, format="JPEG", quality=82, optimize=True)
            return Response(
                content=buffer.getvalue(),
                media_type="image/jpeg",
                headers={"Cache-Control": "public, max-age=86400"},
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to render jiandu thumbnail: {exc}") from exc


def _content_disposition(filename: str) -> str:
    if not filename:
        filename = "document"
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("._")
    if not ascii_name:
        ascii_name = "document"
    quoted = quote(filename)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quoted}"


def _markdown_has_images(markdown: str) -> bool:
    if not markdown:
        return False
    return bool(re.search(r"!\[[^\]]*\]\([^)]+\)", markdown))


def _contains_cjk(text: str) -> bool:
    if not text:
        return False
    return bool(re.search(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]", text))


def _resolve_pdf_cjk_font_name() -> str | None:
    env_name = os.getenv("DEEPBOOK_PDF_CJK_FONT", "").strip()
    if env_name:
        return env_name

    known_fonts = [
        ("Microsoft YaHei", [r"C:\\Windows\\Fonts\\msyh.ttc", r"C:\\Windows\\Fonts\\msyh.ttf"]),
        ("SimSun", [r"C:\\Windows\\Fonts\\simsun.ttc", r"C:\\Windows\\Fonts\\simsun.ttf"]),
        ("Noto Sans CJK SC", []),
        ("PingFang SC", []),
        ("WenQuanYi Zen Hei", []),
    ]
    for font_name, font_paths in known_fonts:
        if not font_paths:
            continue
        if any(Path(path).exists() for path in font_paths):
            return font_name
    return None


def _resolve_public_api_base(request: Request, asset_base: str | None = None) -> str:
    candidate = (asset_base or "").strip()
    if candidate:
        return candidate.rstrip("/")

    configured = os.getenv("DEEPBOOK_PUBLIC_API_BASE", "").strip()
    if configured:
        return configured.rstrip("/")

    origin = request.headers.get("origin", "").strip()
    if origin:
        return f"{origin.rstrip('/')}/api"

    return str(request.base_url).rstrip("/")


def _normalize_markdown_asset_subpath(path_text: str) -> str:
    normalized = (path_text or "").replace("\\", "/")
    while normalized.startswith("./"):
        normalized = normalized[2:]
    while normalized.startswith("../"):
        normalized = normalized[3:]
    normalized = normalized.lstrip("/")
    if not normalized:
        return ""
    return re.sub(
        r"^(?:(?:images|imgs)/)+",
        "images/",
        normalized,
        flags=re.IGNORECASE,
    )


def _rewrite_markdown_asset_url(
    url: str, job_id: str, output_dir_rel: str, public_api_base: str
) -> str:
    if not url or not job_id or not public_api_base:
        return url

    raw = url.strip()
    wrapped = raw.startswith("<") and raw.endswith(">")
    if wrapped:
        raw = raw[1:-1].strip()
    if not raw or raw.startswith("#"):
        return url

    parsed = urlsplit(raw)
    if parsed.scheme or parsed.netloc:
        return url

    rel_path = _normalize_markdown_asset_subpath(parsed.path)
    if not rel_path:
        return url

    base_rel = (output_dir_rel or "").replace("\\", "/").strip("/")
    if base_rel:
        if rel_path.startswith(f"{base_rel}/"):
            suffix = _normalize_markdown_asset_subpath(rel_path[len(base_rel) + 1 :])
            rel_path = f"{base_rel}/{suffix}" if suffix else base_rel
        else:
            rel_path = f"{base_rel}/{rel_path}"

    absolute = (
        f"{public_api_base.rstrip('/')}/ocr/jobs/{job_id}/files/"
        f"{quote(rel_path, safe='/')}"
    )
    if parsed.query:
        absolute = f"{absolute}?{parsed.query}"
    if parsed.fragment:
        absolute = f"{absolute}#{parsed.fragment}"
    if wrapped:
        return f"<{absolute}>"
    return absolute


def _rewrite_markdown_asset_links(
    markdown: str, job_id: str, output_dir_rel: str, public_api_base: str
) -> str:
    if not markdown:
        return markdown

    def replace_markdown_image(match: re.Match[str]) -> str:
        alt_text = match.group(1)
        inner = match.group(2).strip()
        if not inner:
            return match.group(0)

        link_part = inner
        suffix = ""
        if inner.startswith("<"):
            closing_idx = inner.find(">")
            if closing_idx != -1:
                link_part = inner[: closing_idx + 1]
                suffix = inner[closing_idx + 1 :].strip()
        else:
            parts = inner.split(maxsplit=1)
            link_part = parts[0]
            suffix = parts[1] if len(parts) > 1 else ""

        rewritten = _rewrite_markdown_asset_url(
            link_part, job_id, output_dir_rel, public_api_base
        )
        if rewritten == link_part:
            return match.group(0)
        rebuilt = rewritten + (f" {suffix}" if suffix else "")
        return f"![{alt_text}]({rebuilt})"

    output = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", replace_markdown_image, markdown)

    def replace_html_image(match: re.Match[str]) -> str:
        prefix, link, suffix = match.groups()
        rewritten = _rewrite_markdown_asset_url(
            link, job_id, output_dir_rel, public_api_base
        )
        return f"{prefix}{rewritten}{suffix}"

    return re.sub(
        r'(<img\b[^>]*\bsrc=["\'])([^"\']+)(["\'])',
        replace_html_image,
        output,
        flags=re.IGNORECASE,
    )


def _resolve_local_model_path(model_name: str) -> str | None:
    if not model_name:
        return None
    try:
        if Path(model_name).exists():
            return model_name
    except OSError:
        return None
    if "/" not in model_name:
        return None
    try:
        from huggingface_hub import snapshot_download
    except Exception:
        return None
    try:
        return snapshot_download(repo_id=model_name, local_files_only=True)
    except Exception:
        return None


def _export_with_pandoc(
    markdown: str, fmt: str, resource_paths: list[str] | None = None
) -> tuple[bytes | None, str | None]:
    if not markdown:
        return None, "No content to export."
    try:
        import pypandoc
    except Exception as exc:
        return None, f"pypandoc unavailable: {exc}"
    output_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{fmt}") as handle:
            output_path = handle.name
        extra_args: list[str] = []
        if resource_paths:
            safe_paths = [path for path in resource_paths if path]
            if safe_paths:
                extra_args.append(f"--resource-path={os.pathsep.join(safe_paths)}")
        if fmt == "pdf":
            extra_args = ["--pdf-engine=xelatex"]
            extra_args.extend(["-V", "fig-pos=H"])
            if _contains_cjk(markdown):
                cjk_font = _resolve_pdf_cjk_font_name()
                if cjk_font:
                    extra_args.extend(["-V", f"mainfont={cjk_font}"])
                    extra_args.extend(["-V", f"CJKmainfont={cjk_font}"])
            if resource_paths:
                safe_paths = [path for path in resource_paths if path]
                if safe_paths:
                    extra_args.append(f"--resource-path={os.pathsep.join(safe_paths)}")
        pypandoc.convert_text(
            markdown,
            to=fmt,
            format="markdown+raw_tex+link_attributes+raw_html+tex_math_dollars",
            outputfile=output_path,
            extra_args=extra_args,
        )
        return Path(output_path).read_bytes(), None
    except Exception as exc:
        return None, str(exc)
    finally:
        if output_path:
            try:
                os.unlink(output_path)
            except OSError:
                pass


def _short_error_message(message: str | None, limit: int = 360) -> str:
    if not message:
        return ""
    compact = " ".join(message.split())
    if len(compact) <= limit:
        return compact
    return f"{compact[:limit]}..."


def _collect_markdown_image_alts(markdown: str) -> set[str]:
    if not markdown:
        return set()
    alts: set[str] = set()
    for match in re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", markdown):
        alt = re.sub(r"\s+", " ", (match.group(1) or "")).strip()
        if alt:
            alts.add(alt)
    for match in re.finditer(r"<img\b[^>]*>", markdown, re.IGNORECASE):
        alt = re.sub(
            r"\s+", " ", _extract_html_img_attr(match.group(0), "alt")
        ).strip()
        if alt:
            alts.add(alt)
    return alts


def _center_docx_image_alt_paragraphs(docx_bytes: bytes, alt_candidates: set[str]) -> bytes:
    if not docx_bytes:
        return docx_bytes
    try:
        from docx import Document as DocxDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception:
        return docx_bytes

    try:
        buffer = io.BytesIO(docx_bytes)
        doc = DocxDocument(buffer)
    except Exception:
        return docx_bytes

    changed = False
    safe_alts = {alt for alt in alt_candidates if alt and len(alt) <= 80}
    for paragraph in doc.paragraphs:
        raw_text = paragraph.text or ""
        text = re.sub(r"\s+", " ", raw_text).strip()
        style_name = (
            str(getattr(paragraph.style, "name", "") or "").strip().lower()
            if paragraph.style is not None
            else ""
        )
        is_caption_style = "caption" in style_name or "图注" in style_name
        is_alt_line = text in safe_alts
        if not is_caption_style and not is_alt_line:
            continue
        if paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            changed = True

    if not changed:
        return docx_bytes

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()


def _normalize_export_math_delimiters(markdown: str) -> str:
    if not markdown:
        return markdown

    output = markdown
    output = re.sub(
        r"(^|[^\\])\$\$([\s\S]*?)\$\$",
        lambda m: f"{m.group(1)}$${(m.group(2) or '').strip()}$$",
        output,
    )
    output = re.sub(
        r"(^|[^\\])\$([^$\n]+?)\$",
        lambda m: f"{m.group(1)}${(m.group(2) or '').strip()}$",
        output,
    )
    return output


def _normalize_markdown_math_for_md(markdown: str) -> str:
    if not markdown:
        return markdown

    normalized = _normalize_export_math_delimiters(markdown)

    def replace_block_math(match: re.Match[str]) -> str:
        prefix = match.group(1) or ""
        body = (match.group(2) or "").strip()
        if not body:
            return match.group(0)
        return f"{prefix}\\[{body}\\]"

    output = re.sub(
        r"(^|[^\\])\$\$([\s\S]*?)\$\$",
        replace_block_math,
        normalized,
    )

    def replace_inline_math(match: re.Match[str]) -> str:
        prefix = match.group(1) or ""
        body = (match.group(2) or "").strip()
        if not body:
            return match.group(0)
        body = re.sub(r"\s+", " ", body)
        return f"{prefix}\\({body}\\)"

    output = re.sub(
        r"(^|[^\\])\$([^$\n]+?)\$",
        replace_inline_math,
        output,
    )
    output = re.sub(
        r"\\\(\s*([^\n]*?)\s*\\\)",
        lambda m: r"\(" + re.sub(r"\s+", " ", (m.group(1) or "").strip()) + r"\)",
        output,
    )
    output = re.sub(
        r"\\\[\s*([\s\S]*?)\s*\\\]",
        lambda m: r"\[" + (m.group(1) or "").strip() + r"\]",
        output,
    )
    return output


def _normalize_relative_asset_link(raw_link: str) -> str:
    link = (raw_link or "").strip()
    if not link:
        return link
    wrapped = link.startswith("<") and link.endswith(">")
    if wrapped:
        link = link[1:-1].strip()
    if not link or link.startswith("#"):
        return raw_link

    parsed = urlsplit(link)
    if parsed.scheme or parsed.netloc:
        return raw_link

    rel_path = _normalize_markdown_asset_subpath(parsed.path)
    if not rel_path:
        return raw_link

    rebuilt = rel_path
    if parsed.query:
        rebuilt = f"{rebuilt}?{parsed.query}"
    if parsed.fragment:
        rebuilt = f"{rebuilt}#{parsed.fragment}"
    if wrapped:
        return f"<{rebuilt}>"
    return rebuilt


def _extract_html_img_attr(tag: str, attr_name: str) -> str:
    if not tag or not attr_name:
        return ""
    escaped = re.escape(attr_name)
    patterns = [
        rf'\b{escaped}\s*=\s*"([^"]*)"',
        rf"\b{escaped}\s*=\s*'([^']*)'",
        rf"\b{escaped}\s*=\s*“([^”]*)”",
        rf"\b{escaped}\s*=\s*‘([^’]*)’",
        rf"\b{escaped}\s*=\s*([^\s>]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, tag, flags=re.IGNORECASE)
        if not match:
            continue
        value = (match.group(1) or "").strip().strip("\"'“”‘’")
        if value:
            return value
    return ""


def _replace_html_img_src_attr(tag: str, new_src: str) -> str:
    if not tag or not new_src:
        return tag
    patterns = [
        r'(\bsrc\s*=\s*)"[^"]*"',
        r"(\bsrc\s*=\s*)'[^']*'",
        r"(\bsrc\s*=\s*)[“”][^“”]*[“”]",
        r"(\bsrc\s*=\s*)[‘’][^‘’]*[‘’]",
        r"(\bsrc\s*=\s*)([^\s>]+)",
    ]
    for pattern in patterns:
        replaced, count = re.subn(
            pattern,
            lambda m: f'{m.group(1)}"{new_src}"',
            tag,
            count=1,
            flags=re.IGNORECASE,
        )
        if count:
            return replaced
    return tag


def _normalize_image_dimension_value(raw_value: str) -> str:
    value = (raw_value or "").strip().strip("\"'“”‘’")
    if not value:
        return ""
    compact = value.replace(" ", "")
    if re.fullmatch(r"\d+(\.\d+)?%", compact):
        return compact
    if re.fullmatch(r"\d+(\.\d+)?(px|pt|cm|mm|in)", compact, flags=re.IGNORECASE):
        return compact.lower()
    if re.fullmatch(r"\d+(\.\d+)?", compact):
        return f"{compact}px"
    return ""


def _extract_image_width_attr(tag: str) -> str:
    width_from_attr = _normalize_image_dimension_value(_extract_html_img_attr(tag, "width"))
    if width_from_attr:
        return width_from_attr
    style_text = _extract_html_img_attr(tag, "style")
    if not style_text:
        return ""
    match = re.search(r"(?:^|;)\s*width\s*:\s*([^;]+)", style_text, re.IGNORECASE)
    if not match:
        return ""
    return _normalize_image_dimension_value(match.group(1))


def _normalize_markdown_asset_links(markdown: str) -> str:
    if not markdown:
        return markdown

    def replace_markdown_image(match: re.Match[str]) -> str:
        alt_text = match.group(1)
        inner = match.group(2).strip()
        if not inner:
            return match.group(0)

        link_part = inner
        suffix = ""
        if inner.startswith("<"):
            closing_idx = inner.find(">")
            if closing_idx != -1:
                link_part = inner[: closing_idx + 1]
                suffix = inner[closing_idx + 1 :].strip()
        else:
            parts = inner.split(maxsplit=1)
            link_part = parts[0]
            suffix = parts[1] if len(parts) > 1 else ""

        rewritten = _normalize_relative_asset_link(link_part)
        rebuilt = rewritten + (f" {suffix}" if suffix else "")
        return f"![{alt_text}]({rebuilt})"

    output = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", replace_markdown_image, markdown)

    def replace_html_image(match: re.Match[str]) -> str:
        tag = match.group(0)
        raw_src = _extract_html_img_attr(tag, "src")
        if not raw_src:
            return tag
        rewritten = _normalize_relative_asset_link(raw_src)
        return _replace_html_img_src_attr(tag, rewritten)

    return re.sub(
        r"<img\b[^>]*>",
        replace_html_image,
        output,
        flags=re.IGNORECASE,
    )


def _convert_html_images_to_markdown(markdown: str) -> str:
    if not markdown:
        return markdown

    def image_tag_to_markdown(tag: str) -> str:
        raw_src = _extract_html_img_attr(tag, "src")
        if not raw_src:
            return ""
        src = _normalize_relative_asset_link(raw_src)
        alt = re.sub(r"\s+", " ", _extract_html_img_attr(tag, "alt")).strip()
        if not alt or len(alt) > 48 or re.search(r"[<>\"“”]", alt):
            alt = "Image"
        width = _extract_image_width_attr(tag)
        attr_suffix = f"{{width={width}}}" if width else ""
        return f"![{alt}]({src}){attr_suffix}"

    output = re.sub(
        r"<img\b[^>]*>",
        lambda m: image_tag_to_markdown(m.group(0)),
        markdown,
        flags=re.IGNORECASE,
    )
    output = re.sub(
        r"<div\b[^>]*>\s*(!\[[^\]]*\]\([^)]+\)(?:\s*\{[^}]*\})*)\s*</div>",
        r"\1",
        output,
        flags=re.IGNORECASE | re.DOTALL,
    )
    output = re.sub(
        r"<p\b[^>]*>\s*(!\[[^\]]*\]\([^)]+\)(?:\s*\{[^}]*\})*)\s*</p>",
        r"\1",
        output,
        flags=re.IGNORECASE | re.DOTALL,
    )
    return output


def _is_standalone_image_line(line: str) -> bool:
    text = (line or "").strip()
    if not text:
        return False
    if re.fullmatch(r"!\[[^\]]*\]\([^)]+\)(?:\s*\{[^}]*\}\s*)*", text):
        return True
    if re.fullmatch(r"<img\b[^>]*\/?>", text, flags=re.IGNORECASE):
        return True
    if re.fullmatch(
        r"<div\b[^>]*>\s*<img\b[^>]*\/?>\s*</div>",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    ):
        return True
    return False


def _extract_centered_caption_div_text(line: str) -> str:
    text = (line or "").strip()
    if not text:
        return ""
    match = re.match(
        r"^<div\b[^>]*text-align\s*:\s*center[^>]*>([\s\S]*?)</div>$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    inner = (match.group(1) or "").strip()
    if not inner:
        return ""
    if re.search(r"<img\b|<table\b|<tr\b|<td\b", inner, flags=re.IGNORECASE):
        return ""
    plain = re.sub(r"<[^>]+>", " ", inner)
    plain = html.unescape(plain)
    plain = re.sub(r"\s+", " ", plain).strip()
    return plain


def _extract_figure_caption_text(line: str) -> str:
    centered = _extract_centered_caption_div_text(line)
    if centered:
        text = centered
    else:
        text = re.sub(r"\s+", " ", (line or "")).strip()
    if not text or len(text) > 180:
        return ""
    if re.match(r"^(#{1,6}\s+|[-*+]\s+|\d+\.\s+|>|`{3,})", text):
        return ""
    if not centered and re.search(
        r"!\[[^\]]*\]\([^)]+\)|<img\b|<div\b|<p\b|<table\b|<tr\b|<td\b",
        text,
        flags=re.IGNORECASE,
    ):
        return ""
    if re.match(r"^\([A-Za-z0-9ivxIVX]+\)\s+\S+", text):
        return text
    if re.match(r"^(?:fig(?:ure)?\.?\s*\d+|图\s*\d+|表\s*\d+)", text, flags=re.IGNORECASE):
        return text
    return ""


def _center_markdown_figure_captions(markdown: str) -> str:
    if not markdown:
        return markdown

    lines = markdown.replace("\r\n", "\n").split("\n")
    for i, line in enumerate(lines):
        if not _is_standalone_image_line(line):
            continue
        j = i + 1
        while j < len(lines):
            raw = lines[j] if j < len(lines) else ""
            trimmed = raw.strip()
            if not trimmed:
                j += 1
                continue
            # Keep existing centered caption blocks unchanged.
            if _extract_centered_caption_div_text(trimmed):
                j += 1
                continue
            caption_text = _extract_figure_caption_text(trimmed)
            if not caption_text:
                break
            lines[j] = f'<div style="text-align: center;">{html.escape(caption_text, quote=True)}</div>'
            j += 1
    return "\n".join(lines)


def _style_standalone_images_for_markdown(markdown: str, default_width: str = "88%") -> str:
    if not markdown:
        return markdown

    lines = markdown.replace("\r\n", "\n").split("\n")
    image_line_pattern = re.compile(
        r"^[ \t]*!\[([^\]]*)\]\(([^)]+)\)\s*((?:\{[^}]*\}\s*)*)$"
    )
    for index, line in enumerate(lines):
        match = image_line_pattern.match(line or "")
        if not match:
            continue
        alt_text = re.sub(r"\s+", " ", (match.group(1) or "")).strip() or "Image"
        link_part = _extract_markdown_image_link_target(match.group(2) or "")
        if not link_part:
            continue

        attr_text = (match.group(3) or "").strip()
        width_value = ""
        for attr_match in re.finditer(
            r"\bwidth\s*=\s*([^\s{}]+)",
            attr_text,
            flags=re.IGNORECASE,
        ):
            width_value = _normalize_image_dimension_value(attr_match.group(1) or "")
            if width_value:
                break
        if not width_value:
            width_value = _normalize_image_dimension_value(default_width) or "88%"

        safe_src = html.escape(link_part, quote=True)
        safe_alt = html.escape(alt_text, quote=True)
        safe_width = html.escape(width_value, quote=True)
        lines[index] = (
            f'<div style="text-align: center;"><img src="{safe_src}" '
            f'alt="{safe_alt}" width="{safe_width}" /></div>'
        )
    return "\n".join(lines)


def _remove_markdown_image_captions(markdown: str) -> str:
    if not markdown:
        return markdown

    def normalize_attr_part(raw_attr: str) -> str:
        attr_text = (raw_attr or "").strip()
        if not attr_text:
            return ""

        width_value = ""
        height_value = ""
        for match in re.finditer(
            r"\b(width|height)\s*=\s*([^\s{}]+)",
            attr_text,
            flags=re.IGNORECASE,
        ):
            key = (match.group(1) or "").strip().lower()
            value = _normalize_image_dimension_value(match.group(2) or "")
            if not value:
                continue
            if key == "width":
                width_value = value
            elif key == "height":
                height_value = value

        attrs: list[str] = []
        if width_value:
            attrs.append(f"width={width_value}")
        if height_value:
            attrs.append(f"height={height_value}")
        return f"{{{' '.join(attrs)}}}" if attrs else ""

    def replace_markdown_image(match: re.Match[str]) -> str:
        link_part = (match.group(1) or "").strip()
        attr_part = normalize_attr_part(match.group(2) or "")
        suffix = attr_part if attr_part else ""
        return f"![]({link_part}){suffix}"

    return re.sub(
        r"!\[[^\]]*\]\(([^)]+)\)\s*((?:\{[^}]*\}\s*)*)",
        replace_markdown_image,
        markdown,
    )


def _extract_markdown_image_links(line: str) -> list[str]:
    if not line:
        return []
    links: list[str] = []
    for match in re.finditer(r"!\[[^\]]*\]\(([^)]+)\)", line):
        inner = match.group(1).strip()
        if not inner:
            continue
        link_part = inner
        if inner.startswith("<"):
            closing_idx = inner.find(">")
            if closing_idx != -1:
                link_part = inner[: closing_idx + 1]
        else:
            link_part = inner.split(maxsplit=1)[0]
        if link_part:
            links.append(link_part.strip())

    for match in re.finditer(r"<img\b[^>]*>", line, re.IGNORECASE):
        src = _extract_html_img_attr(match.group(0), "src").strip()
        if src:
            links.append(src)
    return links


def _extract_markdown_image_link_target(raw_inner: str) -> str:
    inner = (raw_inner or "").strip()
    if not inner:
        return ""
    if inner.startswith("<"):
        closing_idx = inner.find(">")
        if closing_idx != -1:
            return inner[1:closing_idx].strip()
    return inner.split(maxsplit=1)[0].strip()


def _dimension_to_latex_graphics_option(name: str, raw_value: str) -> str:
    normalized = _normalize_image_dimension_value(raw_value)
    if not normalized:
        return ""
    lower = normalized.lower()
    if lower.endswith("%"):
        try:
            ratio = float(lower[:-1]) / 100.0
        except ValueError:
            return ""
        if ratio <= 0:
            return ""
        return f"{name}={ratio:.4f}\\linewidth"
    if lower.endswith("px"):
        try:
            points = float(lower[:-2]) * 0.75
        except ValueError:
            return ""
        if points <= 0:
            return ""
        return f"{name}={points:.2f}pt"
    return f"{name}={lower}"


def _escape_latex_text(text: str) -> str:
    escaped = (text or "").strip()
    if not escaped:
        return ""
    replacements = {
        "\\": r"\textbackslash{}",
        "{": r"\{",
        "}": r"\}",
        "$": r"\$",
        "&": r"\&",
        "#": r"\#",
        "%": r"\%",
        "_": r"\_",
        "^": r"\^{}",
        "~": r"\~{}",
    }
    return "".join(replacements.get(ch, ch) for ch in escaped)


def _is_likely_figure_caption_line(line: str) -> bool:
    return bool(_extract_figure_caption_text(line))


def _convert_standalone_markdown_images_to_latex(markdown: str) -> str:
    if not markdown:
        return markdown

    lines = markdown.splitlines()
    output_lines: list[str] = []
    i = 0
    image_line_pattern = re.compile(
        r"^[ \t]*!\[[^\]]*\]\(([^)]+)\)\s*((?:\{[^}]*\}\s*)*)$"
    )
    while i < len(lines):
        line = lines[i]
        match = image_line_pattern.match(line or "")
        if not match:
            output_lines.append(line)
            i += 1
            continue

        link_part = _extract_markdown_image_link_target(match.group(1) or "")
        if not link_part:
            output_lines.append(line)
            i += 1
            continue

        attr_text = (match.group(2) or "").strip()
        width_value = ""
        height_value = ""
        for attr_match in re.finditer(
            r"\b(width|height)\s*=\s*([^\s{}]+)",
            attr_text,
            flags=re.IGNORECASE,
        ):
            key = (attr_match.group(1) or "").strip().lower()
            value = attr_match.group(2) or ""
            if key == "width":
                width_value = value
            elif key == "height":
                height_value = value

        options: list[str] = []
        width_opt = _dimension_to_latex_graphics_option("width", width_value)
        if width_opt:
            options.append(width_opt)
        height_opt = _dimension_to_latex_graphics_option("height", height_value)
        if height_opt:
            options.append(height_opt)
        option_suffix = f"[{','.join(options)}]" if options else ""

        caption_text = ""
        look_ahead = i + 1
        skipped_blank = 0
        while look_ahead < len(lines) and not lines[look_ahead].strip() and skipped_blank < 2:
            skipped_blank += 1
            look_ahead += 1
        if look_ahead < len(lines):
            candidate = lines[look_ahead].strip()
            extracted_caption = _extract_figure_caption_text(candidate)
            if extracted_caption:
                caption_text = _escape_latex_text(extracted_caption)

        normalized_path = link_part.replace("\\", "/")
        block_lines = [
            r"\begin{center}",
            f"\\includegraphics{option_suffix}" + f"{{\\detokenize{{{normalized_path}}}}}",
        ]
        if caption_text:
            block_lines.append(r"\par\smallskip")
            block_lines.append(rf"\small {caption_text}")
        block_lines.append(r"\end{center}")
        output_lines.append("\n".join(block_lines))

        if caption_text:
            i = look_ahead + 1
        else:
            i += 1

    return "\n".join(output_lines)


def _strip_markdown_image_markup(line: str) -> str:
    if not line:
        return line
    output = re.sub(r"!\[[^\]]*\]\(([^)]+)\)(?:\s*\{[^}]*\})*", "", line)
    output = re.sub(r"<img\b[^>]*>", "", output, flags=re.IGNORECASE)
    output = re.sub(
        r"\s*\{(?:\s*(?:width|height)\s*=\s*[^\s{}]+\s*)+\}",
        "",
        output,
        flags=re.IGNORECASE,
    )
    return output


def _resolve_export_asset_path(
    raw_link: str, resource_paths: list[str] | None = None
) -> str | None:
    if not raw_link:
        return None
    link = raw_link.strip()
    wrapped = link.startswith("<") and link.endswith(">")
    if wrapped:
        link = link[1:-1].strip()
    if not link:
        return None

    if re.match(r"^[A-Za-z]:[\\/]", link):
        if Path(link).is_file():
            return link
        return None

    parsed = urlsplit(link)
    scheme = parsed.scheme.lower()
    if scheme == "file":
        file_path = unquote(parsed.path or "")
        if re.match(r"^/[A-Za-z]:[\\/]", file_path):
            file_path = file_path.lstrip("/")
        candidate = Path(file_path)
        if candidate.is_file():
            return str(candidate)
        return None

    if parsed.scheme or parsed.netloc:
        return None

    rel_path = unquote(parsed.path or "").replace("\\", "/")
    while rel_path.startswith("./"):
        rel_path = rel_path[2:]
    while rel_path.startswith("../"):
        rel_path = rel_path[3:]
    rel_path = rel_path.lstrip("/")
    if not rel_path:
        return None

    search_paths = [path for path in (resource_paths or []) if path]
    search_paths.append(str(PROJECT_ROOT))
    for base in search_paths:
        candidate = Path(base) / rel_path
        try:
            if candidate.is_file():
                return str(candidate)
        except OSError:
            continue
    return None


def _export_pdf_with_reportlab(
    title: str,
    markdown_content: str,
    text_content: str,
    resource_paths: list[str] | None = None,
    pandoc_error: str | None = None,
) -> Response:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Image as RLImage
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception as exc:
        detail = "reportlab not installed. Install reportlab to export pdf."
        extra = _short_error_message(pandoc_error)
        if extra:
            detail = (
                f"{detail} Pandoc conversion also failed: {extra}"
            )
        raise HTTPException(status_code=500, detail=detail) from exc

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    style = styles["BodyText"]
    style.wordWrap = "CJK"
    style.leading = 18

    font_candidates: list[tuple[str, str]] = [
        ("DocFontYaHeiTTC", r"C:\\Windows\\Fonts\\msyh.ttc"),
        ("DocFontYaHeiTTF", r"C:\\Windows\\Fonts\\msyh.ttf"),
        ("DocFontSimSunTTC", r"C:\\Windows\\Fonts\\simsun.ttc"),
        ("DocFontSimSunTTF", r"C:\\Windows\\Fonts\\simsun.ttf"),
        ("DocFontSimHeiTTF", r"C:\\Windows\\Fonts\\simhei.ttf"),
    ]
    font_loaded = False
    for font_name, font_path in font_candidates:
        if Path(font_path).exists():
            try:
                if font_path.lower().endswith(".ttc"):
                    loaded_ttc = False
                    for sub_index in (0, 1):
                        try:
                            pdfmetrics.registerFont(
                                TTFont(font_name, font_path, subfontIndex=sub_index)
                            )
                            style.fontName = font_name
                            loaded_ttc = True
                            font_loaded = True
                            break
                        except Exception:
                            continue
                    if loaded_ttc:
                        break
                    continue
                pdfmetrics.registerFont(TTFont(font_name, font_path))
                style.fontName = font_name
                font_loaded = True
                break
            except Exception:
                continue
    if not font_loaded:
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            style.fontName = "STSong-Light"
        except Exception:
            pass

    markdown_lines = markdown_content.splitlines()
    source_lines = markdown_lines if markdown_lines else text_content.splitlines()
    story: list[Any] = []

    for raw_line in source_lines:
        line = raw_line.rstrip()
        if not line.strip():
            story.append(Spacer(1, 12))
            continue

        image_links = _extract_markdown_image_links(line)
        remaining_text = _strip_markdown_image_markup(line).strip()
        if remaining_text:
            safe_line = (
                remaining_text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            story.append(Paragraph(safe_line, style))
            story.append(Spacer(1, 8))

        for raw_link in image_links:
            local_path = _resolve_export_asset_path(raw_link, resource_paths)
            if local_path:
                try:
                    image = RLImage(local_path)
                    image.hAlign = "CENTER"
                    image_width = float(image.imageWidth or 0)
                    image_height = float(image.imageHeight or 0)
                    if image_width > 0 and image_height > 0:
                        max_width = float(doc.width)
                        max_height = float(A4[1] * 0.55)
                        scale = min(max_width / image_width, max_height / image_height, 1.0)
                        image.drawWidth = image_width * scale
                        image.drawHeight = image_height * scale
                    story.append(image)
                except Exception:
                    safe_link = (
                        raw_link.replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;")
                    )
                    story.append(Paragraph(f"[图片加载失败: {safe_link}]", style))
            else:
                safe_link = (
                    raw_link.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )
                story.append(Paragraph(f"[图片未找到: {safe_link}]", style))
            story.append(Spacer(1, 10))

    if not story:
        story.append(Paragraph(" ", style))

    doc.build(story)
    buffer.seek(0)
    filename = f"{title}.pdf"
    return Response(
        content=buffer.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


def _pdf_response(title: str, pdf_bytes: bytes) -> Response:
    filename = f"{title}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


def _merge_pdf_chunks(pdf_chunks: list[bytes]) -> bytes:
    valid_chunks = [chunk for chunk in pdf_chunks if chunk]
    if not valid_chunks:
        return b""

    try:
        import fitz  # PyMuPDF
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="pymupdf not installed. Install pymupdf to merge batch PDF exports.",
        ) from exc

    merged = fitz.open()
    try:
        for chunk in valid_chunks:
            src = fitz.open(stream=chunk, filetype="pdf")
            try:
                merged.insert_pdf(src)
            finally:
                src.close()
        return merged.tobytes()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to merge PDF pages: {exc}") from exc
    finally:
        merged.close()


def _index_document(job_id: str, job_result: dict[str, Any]) -> None:
    global FTS_ENABLED
    title = job_result.get("file") or job_result.get("file_name") or ""
    content = (
        job_result.get("raw_text")
        or job_result.get("text")
        or job_result.get("markdown")
        or ""
    )
    output_files = job_result.get("output_files") or []
    cropped_images = job_result.get("cropped_images") or []
    created_at = datetime.now(tz=timezone.utc).isoformat()

    try:
        cursor = _db_execute(
            """
            INSERT INTO documents (
                job_id, title, file_name, markdown, text, raw_text,
                edited_markdown, edited_text, output_dir, output_dir_rel, output_files,
                boxes_image, cropped_images, ocr_model, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                job_id,
                title,
                job_result.get("file"),
                job_result.get("markdown"),
                job_result.get("text"),
                job_result.get("raw_text"),
                None,
                None,
                job_result.get("output_dir"),
                job_result.get("output_dir_rel"),
                json.dumps(output_files, ensure_ascii=False),
                job_result.get("boxes_image"),
                json.dumps(cropped_images, ensure_ascii=False),
                _resolve_ocr_model(str(job_result.get("ocr_model", ""))),
                created_at,
                created_at,
            ),
        )
        doc_id = cursor.lastrowid
        _upsert_document_fts(doc_id, title, content)
        return doc_id
    except Exception as exc:
        print(f"[db] Failed to index document: {exc}")
        return None


def _upsert_document_fts(doc_id: int, title: str, content: str) -> None:
    global FTS_ENABLED
    if not FTS_ENABLED:
        return
    try:
        _db_execute("DELETE FROM documents_fts WHERE rowid = ?", (doc_id,))
        _db_execute(
            "INSERT INTO documents_fts (rowid, title, content) VALUES (?, ?, ?)",
            (doc_id, title, content),
        )
    except sqlite3.Error:
        FTS_ENABLED = False


def _row_to_document(row: sqlite3.Row) -> dict[str, Any]:
    ocr_model = (
        _resolve_ocr_model(str(row["ocr_model"] or ""))
        if "ocr_model" in row.keys()
        else OCR_MODEL_DEEPSEEK
    )
    return {
        "doc_id": row["id"],
        "job_id": row["job_id"],
        "file": row["file_name"] or row["title"] or "",
        "title": row["title"],
        "output_dir": row["output_dir"],
        "output_dir_rel": row["output_dir_rel"],
        "output_files": _deserialize_json(row["output_files"], []),
        "markdown": row["markdown"],
        "text": row["text"],
        "raw_text": row["raw_text"],
        "edited_markdown": row["edited_markdown"],
        "edited_text": row["edited_text"],
        "boxes_image": row["boxes_image"],
        "cropped_images": _deserialize_json(row["cropped_images"], []),
        "ocr_model": ocr_model,
    }


def _hydrate_results_from_db(job_id: str, results: list[dict[str, Any]]) -> None:
    if not results:
        return
    needs_lookup = any(not result.get("doc_id") for result in results)
    if not needs_lookup:
        return
    try:
        rows = _db_query(
            """
            SELECT id, file_name, title, edited_markdown, edited_text
            FROM documents
            WHERE job_id = ?
            """,
            (job_id,),
        )
    except Exception:
        return
    if not rows:
        return
    by_id: dict[int, sqlite3.Row] = {}
    by_file: dict[str, sqlite3.Row] = {}
    for row in rows:
        by_id[row["id"]] = row
        key = row["file_name"] or row["title"] or ""
        if key:
            by_file[key] = row

    for result in results:
        row = None
        doc_id = result.get("doc_id")
        if doc_id in by_id:
            row = by_id[doc_id]
        if row is None:
            name = result.get("file") or result.get("file_name") or ""
            row = by_file.get(name)
        if row is None:
            continue
        result["doc_id"] = row["id"]
        if row["edited_markdown"] is not None:
            result["edited_markdown"] = row["edited_markdown"]
        if row["edited_text"] is not None:
            result["edited_text"] = row["edited_text"]


def _build_snippet(content: str, query: str, radius: int = 60) -> str:
    if not content:
        return ""
    lower = content.lower()
    needle = query.lower()
    idx = lower.find(needle)
    if idx == -1:
        return content[: radius * 2] + ("..." if len(content) > radius * 2 else "")
    start = max(0, idx - radius)
    end = min(len(content), idx + len(query) + radius)
    snippet = content[start:end]
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(content) else ""
    return f"{prefix}{snippet}{suffix}"


def _count_substring_occurrences(text: str, needle: str) -> int:
    if not text or not needle:
        return 0
    count = 0
    start = 0
    step = max(1, len(needle))
    while True:
        idx = text.find(needle, start)
        if idx == -1:
            break
        count += 1
        start = idx + step
    return count


def _tokenize_query(query: str) -> list[str]:
    lowered = query.lower().strip()
    if not lowered:
        return []

    tokens: list[str] = []
    for chunk in re.split(r"\s+", lowered):
        if chunk and chunk not in tokens:
            tokens.append(chunk)

    for chunk in re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]+", lowered):
        if chunk and chunk not in tokens:
            tokens.append(chunk)

    return tokens


def _fallback_relevance_score(
    query: str, title: str, file_name: str, content: str
) -> float:
    q = query.lower().strip()
    if not q:
        return 0.0

    title_text = f"{title or ''} {file_name or ''}".lower()
    body_text = (content or "").lower()
    tokens = _tokenize_query(q)

    phrase_hits_title = _count_substring_occurrences(title_text, q)
    phrase_hits_body = _count_substring_occurrences(body_text, q)
    token_hits_title = sum(
        min(3, _count_substring_occurrences(title_text, token)) for token in tokens
    )
    token_hits_body = sum(
        min(10, _count_substring_occurrences(body_text, token)) for token in tokens
    )

    # Weighted term-frequency score; title/file hits are treated as stronger relevance.
    raw_score = (
        phrase_hits_title * 8
        + phrase_hits_body * 5
        + token_hits_title * 3
        + token_hits_body
    )
    length_penalty = max(1.0, len(content) / 1500.0)
    return raw_score / length_penalty


def _search_documents_fallback(
    query: str, limit: int, candidate_limit: int | None = None
) -> list[dict[str, Any]]:
    like = f"%{query}%"
    safe_candidates = (
        max(limit, candidate_limit) if candidate_limit is not None else max(limit * 20, 1000)
    )
    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name,
            markdown, text, raw_text,
            edited_markdown, edited_text,
            created_at
        FROM documents
        WHERE title LIKE ?
           OR file_name LIKE ?
           OR edited_text LIKE ?
           OR edited_markdown LIKE ?
           OR text LIKE ?
           OR raw_text LIKE ?
           OR markdown LIKE ?
        ORDER BY id DESC
        LIMIT ?
        """,
        (like, like, like, like, like, like, like, safe_candidates),
    )
    results: list[dict[str, Any]] = []
    for row in rows:
        content = (
            row["edited_text"]
            or row["edited_markdown"]
            or row["text"]
            or row["raw_text"]
            or row["markdown"]
            or ""
        )
        score = _fallback_relevance_score(
            query, row["title"] or "", row["file_name"] or "", content
        )
        results.append(
            {
                "id": row["id"],
                "job_id": row["job_id"],
                "title": row["title"],
                "file_name": row["file_name"],
                "created_at": row["created_at"],
                "snippet": _build_snippet(content, query),
                "score": score,
            }
        )
    results.sort(key=lambda item: (item["score"], item["id"]), reverse=True)
    return results[:limit]


def _search_documents(query: str, limit: int) -> list[dict[str, Any]]:
    if not FTS_ENABLED:
        return _search_documents_fallback(query, limit)

    fts_rows = _db_query(
        """
        SELECT
            d.id,
            d.job_id,
            d.title,
            d.file_name,
            d.created_at,
            snippet(documents_fts, 1, '[', ']', '...', 10) AS snippet,
            bm25(documents_fts) AS score
        FROM documents_fts
        JOIN documents d ON d.id = documents_fts.rowid
        WHERE documents_fts MATCH ?
        ORDER BY score
        LIMIT ?
        """,
        (query, limit),
    )
    fts_results = [
        {
            "id": row["id"],
            "job_id": row["job_id"],
            "title": row["title"],
            "file_name": row["file_name"],
            "created_at": row["created_at"],
            "snippet": row["snippet"],
            # bm25 lower is better; keep compatible with existing UI by flipping sign.
            "score": -float(row["score"]),
        }
        for row in fts_rows
    ]
    fallback_results = _search_documents_fallback(
        query, max(limit * 5, limit), candidate_limit=max(limit * 20, 1000)
    )
    fts_ids = {item["id"] for item in fts_results}
    merged = fts_results + [item for item in fallback_results if item["id"] not in fts_ids]
    return merged[:limit]


def _resolve_moonshot_api_key() -> str:
    for env_name in ("MOONSHOT_API_KEY", "KIMI_API_KEY"):
        value = os.getenv(env_name, "").strip()
        if value:
            return value
    return ""


def _normalize_assistant_messages(
    raw_messages: list[AssistantChatMessage] | list[dict[str, Any]] | None,
) -> list[dict[str, str]]:
    if not raw_messages:
        return []

    normalized: list[dict[str, str]] = []
    source = raw_messages[-ASSISTANT_MAX_HISTORY_MESSAGES:]
    for raw in source:
        if isinstance(raw, BaseModel):
            role = str(getattr(raw, "role", "") or "").strip().lower()
            content = str(getattr(raw, "content", "") or "")
        else:
            role = str((raw or {}).get("role") or "").strip().lower()
            content = str((raw or {}).get("content") or "")
        if role not in {"system", "user", "assistant"}:
            continue
        compact = re.sub(r"\r\n?", "\n", content).strip()
        if not compact:
            continue
        normalized.append(
            {
                "role": role,
                "content": compact[:ASSISTANT_MAX_MESSAGE_CHARS],
            }
        )
    return normalized


def _get_last_user_message(messages: list[dict[str, str]]) -> str:
    for item in reversed(messages):
        if item.get("role") == "user" and item.get("content"):
            return item["content"]
    return ""


def _assistant_search_documents(query: str, limit: int = ASSISTANT_CONTEXT_DOC_LIMIT) -> list[dict[str, Any]]:
    normalized_query = re.sub(r"\s+", " ", str(query or "")).strip()
    if not normalized_query:
        return []

    results = _search_documents_fallback(
        normalized_query,
        limit,
        candidate_limit=max(300, limit * 80),
    )
    items: list[dict[str, Any]] = []
    for row in results[:limit]:
        items.append(
            {
                "doc_id": row["id"],
                "job_id": row["job_id"],
                "title": row.get("title") or row.get("file_name") or f"文档 {row['id']}",
                "file_name": row.get("file_name") or "",
                "created_at": row.get("created_at") or "",
                "snippet": _short_error_message(row.get("snippet") or "", limit=180),
            }
        )
    return items


def _assistant_search_gallery(
    query: str,
    request: Request,
    limit: int = ASSISTANT_CONTEXT_GALLERY_LIMIT,
) -> list[dict[str, Any]]:
    normalized_query = re.sub(r"\s+", " ", str(query or "")).strip()
    if not normalized_query:
        return []

    tokens = _tokenize_query(normalized_query)
    records, _, _ = _load_jiandu_gallery()
    scored: list[tuple[float, dict[str, Any]]] = []
    for record in records:
        title = str(record.get("title") or "")
        identity = f"{record.get('id') or ''} {record.get('image_file') or ''}"
        content = str(record.get("reference_text") or "")
        score = _fallback_relevance_score(normalized_query, title, identity, content)
        if score <= 0:
            haystacks = [title.lower(), identity.lower(), content.lower()]
            if not any(token and any(token in hay for hay in haystacks) for token in tokens):
                continue
            score = 1.0
        scored.append((score, record))

    scored.sort(
        key=lambda item: (
            item[0],
            int(item[1].get("num_chars") or 0),
        ),
        reverse=True,
    )
    items: list[dict[str, Any]] = []
    for _, record in scored[:limit]:
        serialized = _serialize_jiandu_gallery_item(record, request, include_boxes=False)
        items.append(
            {
                "id": serialized["id"],
                "title": serialized["title"],
                "reference_text": _short_error_message(serialized.get("reference_text") or "", limit=120),
                "box_count": serialized.get("box_count") or 0,
                "num_chars": serialized.get("num_chars") or 0,
                "image_url": serialized.get("image_url") or "",
            }
        )
    return items


def _build_assistant_platform_context(
    query: str,
    request: Request,
) -> tuple[str, dict[str, Any]]:
    docs = _assistant_search_documents(query, ASSISTANT_CONTEXT_DOC_LIMIT)
    gallery_items = _assistant_search_gallery(query, request, ASSISTANT_CONTEXT_GALLERY_LIMIT)
    records, _, manifest = _load_jiandu_gallery()
    gallery_count = len(records)
    lines = [
        f"平台名称：{DEEPBOOK_PLATFORM_NAME}。",
        "平台功能摘要：",
        "- 简牍藏馆：展示秦简样例图像，可查看检测框与参考释文。",
        f"- 简牍藏馆当前内置样例约 {gallery_count} 条，数据集清单记录为 {manifest.get('num_records', gallery_count)} 条。",
        f"- 秦简识别：当前平台默认模型为 {OCR_MODEL_DEEPJIANDU}，另支持 {OCR_MODEL_DEEPSEEK} 与 {OCR_MODEL_PADDLE}。",
        "- 全文检索：可检索历史 OCR 文档内容与任务结果。",
        "- 历史任务：可复查任务详情、输出结果与导出记录。",
    ]

    if docs:
        lines.append("平台文档摘录：")
        for index, item in enumerate(docs, start=1):
            lines.append(
                f"{index}. 文档《{item['title']}》：{item['snippet'] or '无摘要'}"
            )

    if gallery_items:
        lines.append("简牍藏馆样例摘录：")
        for index, item in enumerate(gallery_items, start=1):
            lines.append(
                f"{index}. 样例 {item['id']}《{item['title']}》：参考释文“{item['reference_text'] or '无'}”。"
            )

    return "\n".join(lines), {"documents": docs, "gallery_items": gallery_items}


def _extract_assistant_text(content: Any) -> str:
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, dict):
                if item.get("type") in {"text", "output_text"} and item.get("text"):
                    parts.append(str(item["text"]))
                elif item.get("content"):
                    parts.append(str(item["content"]))
            elif isinstance(item, str):
                parts.append(item)
        return "\n".join(part.strip() for part in parts if part and str(part).strip()).strip()
    if content is None:
        return ""
    return str(content).strip()


def _call_moonshot_chat_completion(
    messages: list[dict[str, str]],
    *,
    temperature: float | None = None,
) -> dict[str, Any]:
    api_key = _resolve_moonshot_api_key()
    if not api_key:
        raise HTTPException(
            status_code=503,
            detail=(
                "秦简问答尚未配置 Moonshot API Key。"
                "请在 DeepBook/.env.local 或运行环境中设置 MOONSHOT_API_KEY。"
            ),
        )

    payload: dict[str, Any] = {
        "model": MOONSHOT_CHAT_MODEL,
        "messages": messages,
        "stream": False,
    }
    if temperature is not None:
        payload["temperature"] = max(0.0, min(float(temperature), 1.5))

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    try:
        response = requests.post(
            MOONSHOT_CHAT_COMPLETIONS_URL,
            json=payload,
            headers=headers,
            timeout=MOONSHOT_TIMEOUT_SECONDS,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Moonshot 请求失败：{exc}") from exc

    if response.status_code != 200:
        detail = _short_error_message(response.text, limit=420)
        raise HTTPException(
            status_code=502,
            detail=(
                f"Moonshot 请求失败（status={response.status_code}）。"
                f"{detail or '上游服务未返回可解析的错误信息。'}"
            ),
        )

    try:
        data = response.json()
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Moonshot 返回的内容不是合法 JSON。") from exc

    choices = data.get("choices") or []
    if not choices:
        raise HTTPException(status_code=502, detail="Moonshot 未返回有效回答。")

    message = choices[0].get("message") or {}
    assistant_text = _extract_assistant_text(message.get("content"))
    if not assistant_text:
        raise HTTPException(status_code=502, detail="Moonshot 返回了空回答。")

    return {
        "assistant": {
            "role": "assistant",
            "content": assistant_text,
        },
        "model": str(data.get("model") or MOONSHOT_CHAT_MODEL),
        "usage": data.get("usage") or {},
    }


def _coalesce_document_content(row: sqlite3.Row | dict[str, Any]) -> str:
    if isinstance(row, sqlite3.Row):
        getter = row.__getitem__
    else:
        getter = row.get
    return (
        getter("edited_text")
        or getter("edited_markdown")
        or getter("text")
        or getter("raw_text")
        or getter("markdown")
        or ""
    )


def _build_plain_preview(content: str, max_chars: int = 140) -> str:
    compact = re.sub(r"\s+", " ", content or "").strip()
    if len(compact) <= max_chars:
        return compact
    return compact[:max_chars].rstrip() + "..."


def _document_matches_filters(
    row: dict[str, Any],
    job_filter: str,
    file_filter: str,
    start_ts: float | None,
    end_ts: float | None,
) -> bool:
    if job_filter and job_filter not in (row.get("job_id") or "").lower():
        return False
    title_text = f"{row.get('title') or ''} {row.get('file_name') or ''}".lower()
    if file_filter and file_filter not in title_text:
        return False
    if start_ts is not None or end_ts is not None:
        created_ts = _parse_iso_timestamp(row.get("created_at"))
        if start_ts is not None and created_ts < start_ts:
            return False
        if end_ts is not None and created_ts > end_ts:
            return False
    return True


def _job_matches_filters(
    job_payload: dict[str, Any],
    keyword: str,
    status_filter: str,
    start_ts: float | None,
    end_ts: float | None,
) -> bool:
    status = (job_payload.get("status") or "").lower()
    if status_filter and status_filter != "all" and status != status_filter:
        return False
    if keyword:
        file_names = " ".join(job_payload.get("files") or []).lower()
        id_text = (job_payload.get("id") or "").lower()
        if keyword not in id_text and keyword not in file_names:
            return False
    if start_ts is not None or end_ts is not None:
        created_ts = _parse_iso_timestamp(job_payload.get("created_at"))
        if start_ts is not None and created_ts < start_ts:
            return False
        if end_ts is not None and created_ts > end_ts:
            return False
    return True


def _job_to_dict(
    job: OCRJob,
    queue_position: int | None = None,
    include_results: bool = True,
    results_count: int | None = None,
) -> dict[str, Any]:
    resolved_results_count = (
        results_count if results_count is not None else len(job.results)
    )
    job_options = job.options or {}
    resolved_model = _resolve_ocr_model(str(job_options.get("ocr_model", "")))
    if job.page_count is not None:
        resolved_page_count = max(int(job.page_count), resolved_results_count)
    elif job.files:
        resolved_page_count = max(len(job.files), resolved_results_count)
    else:
        resolved_page_count = resolved_results_count
    payload = {
        "id": job.id,
        "status": job.status,
        "progress": job.progress,
        "created_at": datetime.fromtimestamp(job.created_at, tz=timezone.utc).isoformat(),
        "updated_at": datetime.fromtimestamp(job.updated_at, tz=timezone.utc).isoformat(),
        "files": job.files,
        "error": job.error,
        "prompt": job.prompt,
        "options": job.options,
        "ocr_model": resolved_model,
        "pause_requested": job.pause_requested,
        "page_count": resolved_page_count,
        "text_excerpt": "",
    }
    if include_results:
        payload["results"] = job.results
    payload["results_count"] = resolved_results_count
    if queue_position is not None:
        payload["queue_position"] = queue_position
    return payload


def _flatten_preview_text(raw_value: Any) -> str:
    if raw_value is None:
        return ""
    if isinstance(raw_value, str):
        text = raw_value
    else:
        try:
            text = json.dumps(raw_value, ensure_ascii=False)
        except TypeError:
            text = str(raw_value)
    return re.sub(r"\s+", " ", text).strip()


def _record_value(record: sqlite3.Row | dict[str, Any], key: str) -> Any:
    if isinstance(record, sqlite3.Row):
        return record[key] if key in record.keys() else None
    return record.get(key)


def _truncate_text_excerpt(text: str, max_chars: int = 180) -> str:
    normalized = (text or "").strip()
    if not normalized:
        return ""
    if len(normalized) <= max_chars:
        return normalized
    return f"{normalized[: max_chars - 1]}…"


def _extract_preview_text_from_record(record: sqlite3.Row | dict[str, Any]) -> str:
    edited_text = _record_value(record, "edited_text")
    text = _record_value(record, "text")
    edited_markdown = _record_value(record, "edited_markdown")
    markdown = _record_value(record, "markdown")

    for candidate in (edited_text, text):
        normalized = _flatten_preview_text(candidate)
        if normalized:
            return normalized

    markdown_text = _markdown_to_text(edited_markdown) or _markdown_to_text(markdown)
    normalized_markdown = _flatten_preview_text(markdown_text)
    if normalized_markdown:
        return normalized_markdown

    return ""


def _extract_preview_text_from_job(job: OCRJob) -> str:
    for result in job.results:
        preview = _extract_preview_text_from_record(result)
        if preview:
            return preview
    return ""


def _load_job_text_excerpts(job_ids: list[str]) -> dict[str, str]:
    if not job_ids:
        return {}
    unique_ids = [job_id for job_id in dict.fromkeys(job_ids) if job_id]
    if not unique_ids:
        return {}

    excerpts: dict[str, str] = {}
    chunk_size = 300
    for offset in range(0, len(unique_ids), chunk_size):
        chunk = unique_ids[offset : offset + chunk_size]
        placeholders = ",".join("?" for _ in chunk)
        rows = _db_query(
            f"""
            SELECT
                d.job_id,
                d.edited_text,
                d.text,
                d.raw_text,
                d.edited_markdown,
                d.markdown
            FROM documents d
            JOIN (
                SELECT job_id, MIN(id) AS first_id
                FROM documents
                WHERE job_id IN ({placeholders})
                GROUP BY job_id
            ) first_doc
              ON first_doc.job_id = d.job_id
             AND first_doc.first_id = d.id
            """,
            tuple(chunk),
        )

        for row in rows:
            job_id = row["job_id"]
            excerpt = _truncate_text_excerpt(_extract_preview_text_from_record(row))
            if excerpt:
                excerpts[job_id] = excerpt
    return excerpts


def _normalize_archived_input_name(name: str) -> str:
    text = (name or "").strip()
    if not text:
        return ""
    match = re.match(r"^(.+\.pdf)\s+\(page\s+\d+\)$", text, re.IGNORECASE)
    if match:
        return match.group(1).strip()
    return text


def _load_archived_job(job_id: str, include_results: bool = True) -> OCRJob | None:
    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text,
            output_dir, output_dir_rel, output_files, boxes_image,
            cropped_images, ocr_model, created_at, updated_at
        FROM documents
        WHERE job_id = ?
        ORDER BY id
        """,
        (job_id,),
    )
    if not rows:
        return None
    files: list[str] = []
    seen = set()
    for row in rows:
        name = row["file_name"] or row["title"] or ""
        name = _normalize_archived_input_name(name)
        if name and name not in seen:
            files.append(name)
            seen.add(name)
    created_at = min(_parse_iso_timestamp(row["created_at"]) for row in rows)
    updated_at = max(
        _parse_iso_timestamp(row["updated_at"] or row["created_at"]) for row in rows
    )
    archived_model = OCR_MODEL_DEEPSEEK
    for row in rows:
        raw_model = str(row["ocr_model"] or "").strip()
        if not raw_model:
            continue
        archived_model = _resolve_ocr_model(raw_model)
        break
    job = OCRJob(
        id=job_id,
        status="succeeded",
        progress=100,
        created_at=created_at,
        updated_at=updated_at,
        files=files,
        input_paths=[],
        prompt=PROMPT_DEFAULT,
        options={"ocr_model": archived_model},
        pause_requested=False,
        page_count=len(rows),
    )
    if include_results:
        job.results = [_row_to_document(row) for row in rows]
    return job


def _update_job(job_id: str, **kwargs: Any) -> None:
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return
        for key, value in kwargs.items():
            setattr(job, key, value)
        job.updated_at = time.time()


def _register_stream(job_id: str) -> queue.Queue:
    q: queue.Queue = queue.Queue()
    with job_streams_lock:
        job_streams.setdefault(job_id, []).append(q)
    return q


def _unregister_stream(job_id: str, q: queue.Queue) -> None:
    with job_streams_lock:
        streams = job_streams.get(job_id)
        if not streams:
            return
        if q in streams:
            streams.remove(q)
        if not streams:
            job_streams.pop(job_id, None)


def _publish_job_event(job_id: str, event: dict[str, Any]) -> None:
    with job_streams_lock:
        streams = list(job_streams.get(job_id, []))
    if not streams:
        return
    for q in streams:
        try:
            q.put_nowait(event)
        except queue.Full:
            pass


def _wait_for_resume(job_id: str) -> bool:
    while True:
        with jobs_lock:
            job = jobs.get(job_id)
            if not job:
                return False
            if job.status in {"failed", "succeeded"}:
                return False
            pause_requested = job.pause_requested
            if pause_requested:
                if job.status not in {"paused"}:
                    job.status = "paused"
                    job.updated_at = time.time()
            else:
                if job.status in {"paused", "pausing"}:
                    job.status = "running"
                    job.updated_at = time.time()
                return True
        time.sleep(0.5)


def _set_job_paused(job_id: str) -> None:
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return
        job.status = "paused"
        job.pause_requested = False
        job.updated_at = time.time()
        progress = job.progress
    _publish_job_event(
        job_id,
        {"type": "status", "data": {"status": "paused", "progress": progress}},
    )


def _start_queue_worker() -> None:
    global queue_worker_started
    with queue_lock:
        if queue_worker_started:
            return
        queue_worker_started = True
    thread = threading.Thread(target=_queue_worker, daemon=True)
    thread.start()


def _run_queued_job(job_id: str, slot_model: str) -> None:
    try:
        with jobs_lock:
            job = jobs.get(job_id)
        if not job:
            return
        if job.status in {"canceled", "failed", "succeeded", "paused"}:
            return

        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": "running", "progress": job.progress}},
        )
        input_paths = [Path(path) for path in job.input_paths]
        options = job.options or {}
        _run_job(
            job_id,
            input_paths,
            job.prompt,
            int(options.get("base_size", 1024)),
            int(options.get("image_size", 768)),
            bool(options.get("crop_mode", True)),
            _resolve_ocr_model(str(options.get("ocr_model", ""))),
            options,
        )
    finally:
        with queue_lock:
            if running_job_ids.get(slot_model) == job_id:
                running_job_ids[slot_model] = None
            queue_event.set()


def _queue_worker() -> None:
    while True:
        queue_event.wait()
        while True:
            dispatch_job_id: str | None = None
            dispatch_model = OCR_MODEL_DEEPSEEK
            with queue_lock:
                queue_snapshot = list(job_queue)
                running_snapshot = dict(running_job_ids)

            stale_ids: list[str] = []
            candidate_job_id: str | None = None
            candidate_model = OCR_MODEL_DEEPSEEK
            with jobs_lock:
                for queued_id in queue_snapshot:
                    job = jobs.get(queued_id)
                    if job is None or job.status in {
                        "canceled",
                        "failed",
                        "succeeded",
                        "paused",
                    }:
                        stale_ids.append(queued_id)
                        continue
                    if candidate_job_id is not None:
                        continue
                    model = _resolve_ocr_model(str((job.options or {}).get("ocr_model", "")))
                    if running_snapshot.get(model):
                        continue
                    candidate_job_id = queued_id
                    candidate_model = model
                    running_snapshot[model] = queued_id

            with queue_lock:
                if stale_ids:
                    stale = set(stale_ids)
                    job_queue[:] = [queued_id for queued_id in job_queue if queued_id not in stale]

                if (
                    candidate_job_id
                    and candidate_job_id in job_queue
                    and running_job_ids.get(candidate_model) is None
                ):
                    job_queue.remove(candidate_job_id)
                    dispatch_job_id = candidate_job_id
                    dispatch_model = candidate_model
                    running_job_ids[dispatch_model] = dispatch_job_id

                has_queued = bool(job_queue)
                has_running = any(bool(job_id) for job_id in running_job_ids.values())
                if dispatch_job_id is None and (not has_queued or has_running):
                    queue_event.clear()

            if dispatch_job_id is None:
                break

            worker = threading.Thread(
                target=_run_queued_job,
                args=(dispatch_job_id, dispatch_model),
                daemon=True,
            )
            worker.start()


def _enqueue_job(job_id: str) -> None:
    with queue_lock:
        if job_id in job_queue or job_id in running_job_ids.values():
            return
        job_queue.append(job_id)
        queue_event.set()


def _remove_from_queue(job_id: str) -> bool:
    with queue_lock:
        if job_id in job_queue:
            job_queue.remove(job_id)
            if not job_queue:
                queue_event.clear()
            return True
    return False


def _queue_snapshot() -> dict[str, Any]:
    with queue_lock:
        order = list(job_queue)
        running_pairs = [
            (model, job_id) for model, job_id in running_job_ids.items() if job_id
        ]
    with jobs_lock:
        queued_jobs = [
            _job_to_dict(jobs[job_id], queue_position=index + 1, include_results=False)
            for index, job_id in enumerate(order)
            if job_id in jobs
        ]
        running_jobs = [
            _job_to_dict(jobs[job_id], include_results=False)
            for _, job_id in running_pairs
            if job_id in jobs
        ]
        current_job = running_jobs[0] if running_jobs else None
        current_job_id = current_job["id"] if current_job else None
    return {
        "current_job_id": current_job_id,
        "current_job": current_job,
        "running_jobs": running_jobs,
        "queue": queued_jobs,
        "order": order,
    }


def _queue_position(job_id: str) -> int | None:
    with queue_lock:
        try:
            return job_queue.index(job_id) + 1
        except ValueError:
            return None


def _pause_requested(job_id: str) -> bool:
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return False
        return job.pause_requested


def _safe_json(value: Any) -> Any:
    try:
        json.dumps(value)
        return value
    except TypeError:
        return str(value)


def _load_model():
    global MODEL, TOKENIZER
    with model_lock:
        if MODEL is not None and TOKENIZER is not None:
            return MODEL, TOKENIZER

        visible_devices = os.getenv("DEEPSEEK_OCR_VISIBLE_DEVICES")
        if visible_devices:
            os.environ["CUDA_VISIBLE_DEVICES"] = visible_devices

        try:
            import torch
            from transformers import AutoModel, AutoTokenizer
        except Exception as exc:
            raise RuntimeError(
                "DeepSeek OCR dependencies missing. Install torch and transformers."
            ) from exc

        model_name = os.getenv("DEEPSEEK_OCR_MODEL", "deepseek-ai/DeepSeek-OCR-2")
        attn_impl = os.getenv("DEEPSEEK_OCR_ATTN", "flash_attention_2")

        if attn_impl == "flash_attention_2":
            try:
                import flash_attn  # noqa: F401
            except Exception:
                attn_impl = "eager"

        local_only = os.getenv("DEEPSEEK_OCR_LOCAL_ONLY", "1").lower() in {
            "1",
            "true",
            "yes",
        }
        allow_remote = os.getenv("DEEPSEEK_OCR_ALLOW_REMOTE", "").lower() in {
            "1",
            "true",
            "yes",
        }
        if local_only and allow_remote:
            local_only = False

        resolved_model = model_name
        if local_only and not allow_remote:
            local_path = _resolve_local_model_path(model_name)
            if local_path:
                resolved_model = local_path

        try:
            tokenizer = AutoTokenizer.from_pretrained(
                resolved_model, trust_remote_code=True, local_files_only=local_only
            )
            model = AutoModel.from_pretrained(
                resolved_model,
                attn_implementation=attn_impl,
                trust_remote_code=True,
                use_safetensors=True,
                local_files_only=local_only,
            )
        except Exception as exc:
            if allow_remote:
                tokenizer = AutoTokenizer.from_pretrained(
                    model_name, trust_remote_code=True, local_files_only=False
                )
                model = AutoModel.from_pretrained(
                    model_name,
                    attn_implementation=attn_impl,
                    trust_remote_code=True,
                    use_safetensors=True,
                    local_files_only=False,
                )
            else:
                raise RuntimeError(
                    "Local model files not found or incomplete. "
                    f"Tried model={resolved_model}. "
                    "Set DEEPSEEK_OCR_ALLOW_REMOTE=1 to allow download, "
                    "or ensure the model is cached locally and HF_HOME/TRANSFORMERS_CACHE "
                    f"match the cache used for download. Details: {exc}"
                ) from exc

        require_cuda = os.getenv("DEEPSEEK_OCR_REQUIRE_CUDA", "").lower() in {
            "1",
            "true",
            "yes",
        }

        if require_cuda and not torch.cuda.is_available():
            raise RuntimeError("CUDA not available. Check driver and torch install.")

        device_pref = os.getenv("DEEPSEEK_OCR_DEVICE", "auto").lower()
        if device_pref == "cpu":
            device = "cpu"
        elif device_pref == "cuda":
            if not torch.cuda.is_available():
                raise RuntimeError(
                    "CUDA not available. Set DEEPSEEK_OCR_DEVICE=cpu to force CPU."
                )
            device = "cuda"
        else:
            if torch.cuda.is_available():
                device = "cuda"
            elif require_cuda:
                raise RuntimeError("CUDA not available. Check driver and torch install.")
            else:
                device = "cpu"

        dtype = torch.bfloat16 if device == "cuda" else torch.float32
        model = model.eval().to(device).to(dtype)

        MODEL = model
        TOKENIZER = tokenizer
        return MODEL, TOKENIZER


def _read_markdown(output_dir: Path) -> str | None:
    for pattern in ("*.md", "*.mmd", "*.txt"):
        for md_path in output_dir.glob(pattern):
            try:
                return md_path.read_text(encoding="utf-8", errors="ignore")
            except OSError:
                return None
    return None


def _markdown_to_text(markdown: str | None) -> str | None:
    if not markdown:
        return None
    text = re.sub(r"!\[[^\]]*]\([^)]+\)", "", markdown)
    text = re.sub(r"\[([^\]]+)]\([^)]+\)", r"\1", text)
    text = re.sub(r"`{1,3}.*?`{1,3}", "", text, flags=re.DOTALL)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^>\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() or None


def _raw_to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    try:
        return json.dumps(value, ensure_ascii=False, indent=2)
    except TypeError:
        return str(value)


def _relative_to_job(job_id: str, path: Path) -> str | None:
    base_dir = (OUTPUT_ROOT / job_id).resolve()
    try:
        return path.resolve().relative_to(base_dir).as_posix()
    except ValueError:
        return None


def _convert_pdf_to_images(job_id: str, pdf_path: Path) -> list[tuple[Path, str]]:
    try:
        import fitz  # PyMuPDF
    except Exception as exc:
        raise RuntimeError(
            "PDF processing requires PyMuPDF. Install pymupdf and retry."
        ) from exc

    output_dir = INPUT_ROOT / job_id / "pdf_pages" / pdf_path.stem
    output_dir.mkdir(parents=True, exist_ok=True)

    page_entries: list[tuple[Path, str]] = []
    with fitz.open(pdf_path) as doc:
        for index, page in enumerate(doc):
            pix = page.get_pixmap(dpi=200)
            output_path = output_dir / f"{pdf_path.stem}_page_{index + 1:03d}.png"
            pix.save(output_path)
            page_entries.append((output_path, f"{pdf_path.name} (page {index + 1})"))

    return page_entries


def _expand_inputs(job_id: str, input_paths: list[Path]) -> list[tuple[Path, str]]:
    expanded: list[tuple[Path, str]] = []
    for path in input_paths:
        if path.suffix.lower() == ".pdf":
            expanded.extend(_convert_pdf_to_images(job_id, path))
        else:
            expanded.append((path, path.name))
    if not expanded:
        raise RuntimeError("No valid input files to process.")
    return expanded


def _get_ocr_helpers(model) -> tuple[Any, Any, Any] | None:
    global OCR_HELPERS
    if OCR_HELPERS is not None:
        return OCR_HELPERS
    with ocr_helpers_lock:
        if OCR_HELPERS is not None:
            return OCR_HELPERS
        try:
            module = importlib.import_module(model.__class__.__module__)
            re_match = getattr(module, "re_match", None)
            process_image_with_refs = getattr(module, "process_image_with_refs", None)
            load_image = getattr(module, "load_image", None)
            if not all([re_match, process_image_with_refs, load_image]):
                return None
            OCR_HELPERS = (re_match, process_image_with_refs, load_image)
            return OCR_HELPERS
        except Exception:
            return None


def _save_outputs_from_raw(
    raw_output: str, image_path: Path, output_dir: Path, model: Any
) -> None:
    if not raw_output:
        return
    outputs = raw_output.strip()
    stop_str = "<｜end▁of▁sentence｜>"
    if outputs.endswith(stop_str):
        outputs = outputs[: -len(stop_str)]
    outputs = outputs.strip()

    helpers = _get_ocr_helpers(model)
    if helpers:
        re_match, process_image_with_refs, load_image = helpers
        matches_ref, matches_images, matches_other = re_match(outputs)
        image = load_image(str(image_path))
        if image is not None:
            image_draw = image.copy()
            (output_dir / "images").mkdir(parents=True, exist_ok=True)
            try:
                result_image = process_image_with_refs(
                    image_draw, matches_ref, str(output_dir)
                )
                if result_image is not None:
                    result_image.save(output_dir / "result_with_boxes.jpg")
            except Exception:
                pass
        else:
            matches_images, matches_other = [], []
    else:
        matches_images, matches_other = [], []

    for idx, a_match_image in enumerate(matches_images):
        outputs = outputs.replace(a_match_image, f"![](images/{idx}.jpg)\n")

    for a_match_other in matches_other:
        outputs = (
            outputs.replace(a_match_other, "")
            .replace("\\coloneqq", ":=")
            .replace("\\eqqcolon", "=:")
        )

    try:
        (output_dir / "result.mmd").write_text(outputs, encoding="utf-8")
    except OSError:
        pass


def _worker_loop(
    task_queue: multiprocessing.Queue,
    result_queue: multiprocessing.Queue,
    prompt: str,
    base_size: int,
    image_size: int,
    crop_mode: bool,
) -> None:
    try:
        model, tokenizer = _load_model()
    except Exception as exc:
        result_queue.put({"error": f"Model load failed: {exc}"})
        return

    while True:
        task = task_queue.get()
        if task is None:
            break

        file_path = Path(task["file_path"])
        output_dir = Path(task["output_dir"])
        display_name = task["display_name"]

        try:
            output_dir.mkdir(parents=True, exist_ok=True)
            raw_output = None
            try:
                raw_output = model.infer(
                    tokenizer,
                    prompt=prompt,
                    image_file=str(file_path),
                    output_path=str(output_dir),
                    base_size=base_size,
                    image_size=image_size,
                    crop_mode=crop_mode,
                    save_results=False,
                    eval_mode=True,
                )
            except Exception:
                raw_output = None

            if raw_output:
                _save_outputs_from_raw(raw_output, file_path, output_dir, model)
                boxes_path = output_dir / "result_with_boxes.jpg"
                if not boxes_path.exists():
                    try:
                        model.infer(
                            tokenizer,
                            prompt=prompt,
                            image_file=str(file_path),
                            output_path=str(output_dir),
                            base_size=base_size,
                            image_size=image_size,
                            crop_mode=crop_mode,
                            save_results=True,
                        )
                    except Exception:
                        pass
            else:
                model.infer(
                    tokenizer,
                    prompt=prompt,
                    image_file=str(file_path),
                    output_path=str(output_dir),
                    base_size=base_size,
                    image_size=image_size,
                    crop_mode=crop_mode,
                    save_results=True,
                )

            result_queue.put(
                {
                    "display_name": display_name,
                    "output_dir": str(output_dir),
                    "raw_output": raw_output,
                }
            )
        except Exception as exc:
            result_queue.put(
                {
                    "error": str(exc),
                    "display_name": display_name,
                    "output_dir": str(output_dir),
                }
            )


def _run_job_deepseek(
    job_id: str,
    input_paths: list[Path],
    prompt: str,
    base_size: int,
    image_size: int,
    crop_mode: bool,
) -> None:
    _update_job(job_id, status="running", progress=0)
    _publish_job_event(
        job_id, {"type": "status", "data": {"status": "running", "progress": 0}}
    )
    try:
        expanded_inputs = _expand_inputs(job_id, input_paths)
        total = max(len(expanded_inputs), 1)
        _update_job(job_id, page_count=total)

        ctx = multiprocessing.get_context("spawn")
        worker = None
        task_queue: multiprocessing.Queue | None = None
        result_queue: multiprocessing.Queue | None = None

        def start_worker():
            nonlocal worker, task_queue, result_queue
            task_queue = ctx.Queue()
            result_queue = ctx.Queue()
            worker = ctx.Process(
                target=_worker_loop,
                args=(task_queue, result_queue, prompt, base_size, image_size, crop_mode),
                daemon=True,
            )
            worker.start()

        def stop_worker():
            nonlocal worker, task_queue
            if worker is None:
                return
            try:
                if task_queue is not None:
                    task_queue.put(None)
            except Exception:
                pass
            try:
                if worker.is_alive():
                    worker.terminate()
            except Exception:
                pass
            try:
                worker.join(timeout=5)
            except Exception:
                pass
            worker = None

        start_worker()

        with jobs_lock:
            job = jobs.get(job_id)
            start_index = len(job.results) if job else 0

        index = max(0, start_index)
        while index < len(expanded_inputs):
            if _pause_requested(job_id):
                stop_worker()
                _set_job_paused(job_id)
                return

            file_path, display_name = expanded_inputs[index]
            output_dir = OUTPUT_ROOT / job_id / file_path.stem
            output_dir.mkdir(parents=True, exist_ok=True)

            if task_queue is None or result_queue is None:
                raise RuntimeError("OCR worker not available.")

            task_queue.put(
                {
                    "file_path": str(file_path),
                    "display_name": display_name,
                    "output_dir": str(output_dir),
                }
            )

            raw_output = None

            while True:
                if _pause_requested(job_id):
                    stop_worker()
                    shutil.rmtree(output_dir, ignore_errors=True)
                    _set_job_paused(job_id)
                    return
                try:
                    message = result_queue.get(timeout=0.5)
                except queue.Empty:
                    if worker is not None and not worker.is_alive():
                        raise RuntimeError("OCR worker exited unexpectedly.")
                    continue

                if message.get("error"):
                    raise RuntimeError(message["error"])

                raw_output = message.get("raw_output")
                break

            markdown = _read_markdown(output_dir)
            text = _markdown_to_text(markdown)
            raw_text = raw_output or markdown or text

            try:
                output_files = [
                    path.name for path in output_dir.iterdir() if path.is_file()
                ]
            except OSError:
                output_files = []

            boxes_image_path = output_dir / "result_with_boxes.jpg"
            boxes_image = (
                _relative_to_job(job_id, boxes_image_path)
                if boxes_image_path.exists()
                else None
            )

            cropped_dir = output_dir / "images"
            cropped_images: list[str] = []
            if cropped_dir.exists():
                for image_path in sorted(cropped_dir.glob("*")):
                    if image_path.is_file():
                        rel_path = _relative_to_job(job_id, image_path)
                        if rel_path:
                            cropped_images.append(rel_path)

            job_result = {
                "job_id": job_id,
                "file": display_name,
                "output_dir": str(output_dir),
                "output_dir_rel": _relative_to_job(job_id, output_dir),
                "output_files": output_files,
                "markdown": markdown,
                "text": text,
                "raw_text": raw_text,
                "boxes_image": boxes_image,
                "cropped_images": cropped_images,
                "raw": _safe_json(raw_output),
            }

            with jobs_lock:
                job = jobs.get(job_id)
                if job:
                    job.results.append(job_result)

            doc_id = _index_document(job_id, job_result)
            if doc_id:
                job_result["doc_id"] = doc_id

            progress = int(((index + 1) / total) * 100)
            _update_job(job_id, progress=progress)
            _publish_job_event(
                job_id,
                {
                    "type": "result",
                    "data": {
                        "result": job_result,
                        "progress": progress,
                        "status": "running",
                    },
                },
            )
            index += 1

        stop_worker()
        _update_job(job_id, pause_requested=False)
        _update_job(job_id, status="succeeded", progress=100)
        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": "succeeded", "progress": 100}},
        )
    except Exception as exc:
        _update_job(job_id, status="failed", error=str(exc))
        _publish_job_event(
            job_id,
            {
                "type": "status",
                "data": {"status": "failed", "error": str(exc)},
            },
        )


def _sanitize_output_component(value: str, fallback: str = "output") -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip()).strip("._")
    return normalized or fallback


def _guess_image_suffix(path_text: str, url: str) -> str:
    candidates = [
        Path(urlsplit(path_text or "").path).suffix.lower(),
        Path(urlsplit(url or "").path).suffix.lower(),
    ]
    for suffix in candidates:
        if suffix in {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}:
            return suffix
    return ".jpg"


def _build_markdown_image_aliases(source_path: str) -> set[str]:
    source = (source_path or "").replace("\\", "/").strip()
    while source.startswith("./"):
        source = source[2:]
    source = source.lstrip("/")
    if not source:
        return set()

    aliases = {source, Path(source).name}

    # Paddle markdown may use mixed prefixes like images/... or imgs/... and
    # also duplicate them (e.g. images/images/... or imgs/images/...).
    collapsed = re.sub(r"^(?:(?:images|imgs)/)+", "images/", source, flags=re.IGNORECASE)
    if collapsed:
        aliases.add(collapsed)

    if collapsed.startswith("images/"):
        tail = collapsed[len("images/") :]
        if tail:
            aliases.add(f"images/{tail}")
            aliases.add(f"imgs/{tail}")
            aliases.add(f"images/images/{tail}")
            aliases.add(f"imgs/images/{tail}")
            aliases.add(f"images/imgs/{tail}")
    else:
        aliases.add(f"images/{collapsed}")
        aliases.add(f"imgs/{collapsed}")
        aliases.add(f"images/images/{collapsed}")
        aliases.add(f"imgs/images/{collapsed}")

    return {alias for alias in aliases if alias}


def _download_binary(url: str, timeout: int = 60) -> bytes:
    if not url:
        raise RuntimeError("Image URL is empty.")
    response = requests.get(url, timeout=timeout)
    response.raise_for_status()
    return response.content


def _collect_output_artifacts(
    job_id: str, output_dir: Path
) -> tuple[list[str], str | None, list[str]]:
    try:
        output_files = [path.name for path in output_dir.iterdir() if path.is_file()]
    except OSError:
        output_files = []

    boxes_image = None
    for candidate in sorted(output_dir.glob("result_with_boxes*")):
        if candidate.is_file():
            boxes_image = _relative_to_job(job_id, candidate)
            if boxes_image:
                break

    cropped_images: list[str] = []
    cropped_dir = output_dir / "images"
    if cropped_dir.exists():
        for image_path in sorted(cropped_dir.rglob("*")):
            if image_path.is_file():
                rel_path = _relative_to_job(job_id, image_path)
                if rel_path:
                    cropped_images.append(rel_path)

    return output_files, boxes_image, cropped_images


def _append_job_result(
    job_id: str,
    display_name: str,
    output_dir: Path,
    markdown: str | None,
    text: str | None,
    raw_text: str | None,
    raw_payload: Any,
) -> dict[str, Any]:
    output_files, boxes_image, cropped_images = _collect_output_artifacts(job_id, output_dir)
    resolved_model = OCR_MODEL_DEEPSEEK
    with jobs_lock:
        job = jobs.get(job_id)
        if job:
            resolved_model = _resolve_ocr_model(str((job.options or {}).get("ocr_model", "")))
    job_result = {
        "job_id": job_id,
        "file": display_name,
        "output_dir": str(output_dir),
        "output_dir_rel": _relative_to_job(job_id, output_dir),
        "output_files": output_files,
        "markdown": markdown,
        "text": text,
        "raw_text": raw_text,
        "boxes_image": boxes_image,
        "cropped_images": cropped_images,
        "ocr_model": resolved_model,
        "raw": _safe_json(raw_payload),
    }

    with jobs_lock:
        job = jobs.get(job_id)
        if job:
            job.results.append(job_result)
            job.updated_at = time.time()

    doc_id = _index_document(job_id, job_result)
    if doc_id:
        job_result["doc_id"] = doc_id
    return job_result


def _resolve_paddle_token(job_options: dict[str, Any]) -> str:
    explicit = str(job_options.get("paddle_token", "")).strip()
    if explicit:
        return explicit
    env_token = (
        os.getenv("PADDLEOCR_VL_TOKEN", "").strip()
        or os.getenv("PADDLE_OCR_VL_TOKEN", "").strip()
        or os.getenv("PADDLE_OCR_TOKEN", "").strip()
    )
    if env_token:
        return env_token
    return "548a3b5b26d32525569eb8ab3c416947ae8295b7"


def _request_paddle_layout(file_path: Path, job_options: dict[str, Any]) -> list[dict[str, Any]]:
    api_url = os.getenv(
        "PADDLEOCR_VL_SYNC_URL",
        "https://g38bkakem3ubo6wc.aistudio-app.com/layout-parsing",
    ).strip()
    token = _resolve_paddle_token(job_options)
    timeout = int(os.getenv("PADDLEOCR_VL_TIMEOUT", "180"))

    file_type = 0 if file_path.suffix.lower() == ".pdf" else 1
    try:
        file_data = base64.b64encode(file_path.read_bytes()).decode("ascii")
    except OSError as exc:
        raise RuntimeError(f"Failed to read input file: {file_path}") from exc

    markdown_ignore_labels = _normalize_paddle_markdown_ignore_labels(
        job_options.get("markdown_ignore_labels")
    )
    layout_shape_mode = _normalize_paddle_layout_shape_mode(
        job_options.get("layout_shape_mode"), "auto"
    )
    prompt_label = _normalize_paddle_prompt_label(job_options.get("prompt_label"), "ocr")
    layout_merge_bboxes_mode = _normalize_paddle_layout_merge_mode(
        job_options.get("layout_merge_bboxes_mode"), "large"
    )
    layout_threshold = _coerce_float_option(
        job_options.get("layout_threshold"), 0.5, min_value=0.0
    )
    layout_threshold = min(1.0, layout_threshold)
    layout_unclip_ratio = _coerce_float_option(
        job_options.get("layout_unclip_ratio"), 1.0, min_value=0.000001
    )
    min_pixels = _coerce_int_option(job_options.get("min_pixels"), 147384, min_value=1)
    max_pixels = _coerce_int_option(job_options.get("max_pixels"), 2822400, min_value=1)
    if min_pixels > max_pixels:
        min_pixels, max_pixels = max_pixels, min_pixels
    payload = {
        "file": file_data,
        "fileType": file_type,
        "markdownIgnoreLabels": markdown_ignore_labels,
        "useDocOrientationClassify": _coerce_bool_option(
            job_options.get("use_doc_orientation_classify"), False
        ),
        "useDocUnwarping": _coerce_bool_option(job_options.get("use_doc_unwarping"), False),
        "useLayoutDetection": _coerce_bool_option(
            job_options.get("use_layout_detection"), True
        ),
        "useChartRecognition": _coerce_bool_option(
            job_options.get("use_chart_recognition"), False
        ),
        "layoutThreshold": layout_threshold,
        "layoutNms": _coerce_bool_option(job_options.get("layout_nms"), True),
        "layoutUnclipRatio": layout_unclip_ratio,
        "layoutMergeBboxesMode": layout_merge_bboxes_mode,
        "mergeTables": _coerce_bool_option(job_options.get("merge_tables"), True),
        "relevelTitles": _coerce_bool_option(job_options.get("relevel_titles"), True),
        "layoutShapeMode": layout_shape_mode,
        "promptLabel": prompt_label,
        "repetitionPenalty": _coerce_float_option(
            job_options.get("repetition_penalty"), 1.0, min_value=0.0
        ),
        "temperature": _coerce_float_option(job_options.get("temperature"), 0.0, min_value=0.0),
        "topP": _coerce_float_option(job_options.get("top_p"), 1.0, min_value=0.0),
        "minPixels": min_pixels,
        "maxPixels": max_pixels,
        "showFormulaNumber": _coerce_bool_option(
            job_options.get("show_formula_number"), False
        ),
        "restructurePages": _coerce_bool_option(job_options.get("restructure_pages"), False),
        "prettifyMarkdown": _coerce_bool_option(job_options.get("prettify_markdown"), False),
        "visualize": _coerce_bool_option(job_options.get("visualize"), True),
    }
    headers = {
        "Authorization": f"token {token}",
        "Content-Type": "application/json",
    }

    try:
        response = requests.post(api_url, json=payload, headers=headers, timeout=timeout)
    except requests.RequestException as exc:
        raise RuntimeError(f"PaddleOCR request failed: {exc}") from exc

    if response.status_code != 200:
        detail = " ".join((response.text or "").split())
        detail = detail[:320] + ("..." if len(detail) > 320 else "")
        raise RuntimeError(
            "PaddleOCR request failed "
            f"(status={response.status_code}). Response: {detail}"
        )

    try:
        data = response.json()
    except Exception as exc:
        raise RuntimeError("PaddleOCR response is not valid JSON.") from exc

    result = data.get("result") or {}
    layout_results = result.get("layoutParsingResults") or []
    if not isinstance(layout_results, list) or not layout_results:
        raise RuntimeError("PaddleOCR returned empty layoutParsingResults.")
    return layout_results


def _save_paddle_page(
    job_id: str,
    file_path: Path,
    display_name: str,
    page_data: dict[str, Any],
    page_index: int,
    total_pages: int,
) -> dict[str, Any]:
    is_pdf = file_path.suffix.lower() == ".pdf"
    output_base = _sanitize_output_component(file_path.stem, "document")
    if total_pages > 1 or is_pdf:
        output_name = f"{output_base}_page_{page_index:03d}"
        page_display_name = f"{display_name} (page {page_index})"
    else:
        output_name = output_base
        page_display_name = display_name

    output_dir = OUTPUT_ROOT / job_id / output_name
    if output_dir.exists():
        output_dir = OUTPUT_ROOT / job_id / f"{output_name}_{uuid.uuid4().hex[:6]}"
    output_dir.mkdir(parents=True, exist_ok=True)

    markdown_payload = page_data.get("markdown") if isinstance(page_data, dict) else {}
    markdown_text = ""
    markdown_images = {}
    if isinstance(markdown_payload, dict):
        markdown_text = str(markdown_payload.get("text") or "")
        markdown_images = markdown_payload.get("images") or {}
    elif markdown_payload:
        markdown_text = str(markdown_payload)

    images_dir = output_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    replace_map: dict[str, str] = {}
    if isinstance(markdown_images, dict):
        for idx, (source_path, image_url) in enumerate(markdown_images.items()):
            source_text = str(source_path or "").replace("\\", "/")
            image_suffix = _guess_image_suffix(source_text, str(image_url or ""))
            file_base = Path(source_text).name or f"image_{idx + 1}"
            file_stem = _sanitize_output_component(Path(file_base).stem, f"image_{idx + 1}")
            saved_rel = f"images/{file_stem}{image_suffix}"
            dest_path = output_dir / saved_rel
            dest_path.parent.mkdir(parents=True, exist_ok=True)
            try:
                dest_path.write_bytes(_download_binary(str(image_url), timeout=90))
            except Exception:
                continue
            for alias in _build_markdown_image_aliases(source_text):
                replace_map[alias] = saved_rel

    for old_path, new_path in replace_map.items():
        if old_path:
            markdown_text = markdown_text.replace(old_path, new_path)

    output_images = page_data.get("outputImages") if isinstance(page_data, dict) else {}
    if isinstance(output_images, dict):
        for idx, (img_name, img_url) in enumerate(output_images.items()):
            suffix = _guess_image_suffix(str(img_name or ""), str(img_url or ""))
            if idx == 0:
                target_path = output_dir / f"result_with_boxes{suffix}"
            else:
                safe_name = _sanitize_output_component(str(img_name), f"output_{idx + 1}")
                target_path = output_dir / f"{safe_name}_{page_index:03d}{suffix}"
            try:
                target_path.write_bytes(_download_binary(str(img_url), timeout=90))
            except Exception:
                continue

    try:
        (output_dir / "result.mmd").write_text(markdown_text, encoding="utf-8")
    except OSError:
        pass

    text = _markdown_to_text(markdown_text)
    raw_text = text or markdown_text
    return _append_job_result(
        job_id=job_id,
        display_name=page_display_name,
        output_dir=output_dir,
        markdown=markdown_text,
        text=text,
        raw_text=raw_text,
        raw_payload=page_data,
    )


def _run_job_paddle(
    job_id: str,
    input_paths: list[Path],
    job_options: dict[str, Any],
) -> None:
    _update_job(job_id, status="running", progress=0)
    _publish_job_event(
        job_id, {"type": "status", "data": {"status": "running", "progress": 0}}
    )

    try:
        total_files = max(len(input_paths), 1)
        _update_job(job_id, page_count=total_files)

        with jobs_lock:
            job = jobs.get(job_id)
            options = dict(job.options or {}) if job else {}
            start_index = int(options.get("paddle_next_file_index", 0))
            if start_index < 0:
                start_index = 0
        if start_index >= len(input_paths):
            start_index = 0

        for file_index in range(start_index, len(input_paths)):
            if _pause_requested(job_id):
                _set_job_paused(job_id)
                return

            file_path = input_paths[file_index]
            display_name = file_path.name
            layout_pages = _request_paddle_layout(file_path, job_options)
            total_pages = len(layout_pages)
            if total_pages <= 0:
                raise RuntimeError(f"PaddleOCR returned empty result for {display_name}.")

            for page_idx, page_data in enumerate(layout_pages, start=1):
                if _pause_requested(job_id):
                    _set_job_paused(job_id)
                    return

                job_result = _save_paddle_page(
                    job_id=job_id,
                    file_path=file_path,
                    display_name=display_name,
                    page_data=page_data if isinstance(page_data, dict) else {},
                    page_index=page_idx,
                    total_pages=total_pages,
                )
                progress = int(((file_index + 1) / total_files) * 100)
                _update_job(job_id, progress=progress)
                _publish_job_event(
                    job_id,
                    {
                        "type": "result",
                        "data": {
                            "result": job_result,
                            "progress": progress,
                            "status": "running",
                        },
                    },
                )

            with jobs_lock:
                job = jobs.get(job_id)
                if job:
                    options = dict(job.options or {})
                    options["paddle_next_file_index"] = file_index + 1
                    job.options = options
                    job.updated_at = time.time()

        with jobs_lock:
            job = jobs.get(job_id)
            if job:
                options = dict(job.options or {})
                options.pop("paddle_next_file_index", None)
                job.options = options
                job.page_count = len(job.results)
                job.updated_at = time.time()

        _update_job(job_id, pause_requested=False, status="succeeded", progress=100)
        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": "succeeded", "progress": 100}},
        )
    except Exception as exc:
        _update_job(job_id, status="failed", error=str(exc))
        _publish_job_event(
            job_id,
            {
                "type": "status",
                "data": {"status": "failed", "error": str(exc)},
            },
        )


def _save_deepjiandu_char_crops(output_dir: Path, image: Image.Image, chars: list[dict[str, Any]]) -> None:
    images_dir = output_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    for index, char_info in enumerate(chars, start=1):
        box = char_info.get("box_xyxy") or {}
        x1 = max(0, int(box.get("x1", 0)))
        y1 = max(0, int(box.get("y1", 0)))
        x2 = min(image.width, int(box.get("x2", x1)))
        y2 = min(image.height, int(box.get("y2", y1)))
        if x2 <= x1 or y2 <= y1:
            continue
        crop = image.crop((x1, y1, x2, y2))
        crop.save(images_dir / f"char_{index:03d}.png")


def _run_job_deepjiandu(
    job_id: str,
    input_paths: list[Path],
    job_options: dict[str, Any],
) -> None:
    _update_job(job_id, status="running", progress=0)
    _publish_job_event(
        job_id, {"type": "status", "data": {"status": "running", "progress": 0}}
    )
    try:
        expanded_inputs = _expand_inputs(job_id, input_paths)
        total = max(len(expanded_inputs), 1)
        _update_job(job_id, page_count=total)

        with jobs_lock:
            job = jobs.get(job_id)
            start_index = len(job.results) if job else 0

        index = max(0, start_index)
        while index < len(expanded_inputs):
            if _pause_requested(job_id):
                _set_job_paused(job_id)
                return

            file_path, display_name = expanded_inputs[index]
            output_name = _sanitize_output_component(file_path.stem, "document")
            output_dir = OUTPUT_ROOT / job_id / output_name
            shutil.rmtree(output_dir, ignore_errors=True)
            output_dir.mkdir(parents=True, exist_ok=True)

            with Image.open(file_path) as image_file:
                image = image_file.convert("RGB")

            result_payload, overlay = run_deepjiandu_page(image, job_options or {})
            text = str(result_payload.get("text") or "")
            markdown = text
            raw_text = json.dumps(result_payload, ensure_ascii=False, indent=2)

            (output_dir / "result.mmd").write_text(markdown, encoding="utf-8")
            (output_dir / "result.json").write_text(raw_text, encoding="utf-8")
            overlay.save(output_dir / "result_with_boxes.png")
            _save_deepjiandu_char_crops(
                output_dir,
                image,
                list(result_payload.get("chars") or []),
            )

            job_result = _append_job_result(
                job_id=job_id,
                display_name=display_name,
                output_dir=output_dir,
                markdown=markdown,
                text=text,
                raw_text=raw_text,
                raw_payload=result_payload,
            )
            progress = int(((index + 1) / total) * 100)
            _update_job(job_id, progress=progress)
            _publish_job_event(
                job_id,
                {
                    "type": "result",
                    "data": {
                        "result": job_result,
                        "progress": progress,
                        "status": "running",
                    },
                },
            )
            index += 1

        _update_job(job_id, pause_requested=False, status="succeeded", progress=100)
        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": "succeeded", "progress": 100}},
        )
    except Exception as exc:
        _update_job(job_id, status="failed", error=str(exc))
        _publish_job_event(
            job_id,
            {
                "type": "status",
                "data": {"status": "failed", "error": str(exc)},
            },
        )


def _run_job(
    job_id: str,
    input_paths: list[Path],
    prompt: str,
    base_size: int,
    image_size: int,
    crop_mode: bool,
    ocr_model: str,
    job_options: dict[str, Any],
) -> None:
    resolved_model = _resolve_ocr_model(ocr_model, strict=False)
    if resolved_model == OCR_MODEL_PADDLE:
        _run_job_paddle(job_id, input_paths, job_options or {})
        return
    if resolved_model == OCR_MODEL_DEEPJIANDU:
        _run_job_deepjiandu(job_id, input_paths, job_options or {})
        return
    _run_job_deepseek(job_id, input_paths, prompt, base_size, image_size, crop_mode)


@app.get("/")
async def root():
    return {"message": f"{DEEPBOOK_PLATFORM_NAME} API is running"}


@app.get("/jiandu/gallery")
async def list_jiandu_gallery(
    request: Request,
    q: str = "",
    page: int | None = None,
    page_size: int | None = None,
):
    records, _, manifest = _load_jiandu_gallery()
    query = str(q or "").strip()
    if query:
        filtered = [
            record
            for record in records
            if query in record.get("id", "") or query in record.get("reference_text", "")
        ]
    else:
        filtered = records

    safe_page, safe_page_size = _normalize_pagination(
        page,
        page_size,
        default_page_size=36,
        max_page_size=80,
    )
    paged_items, total, total_pages, has_prev, has_next = _paginate_items(
        filtered, safe_page, safe_page_size
    )
    items = [
        _serialize_jiandu_gallery_item(record, request, include_boxes=False)
        for record in paged_items
    ]
    return {
        "items": items,
        "query": query,
        "page": safe_page,
        "page_size": safe_page_size,
        "total": total,
        "total_pages": total_pages,
        "has_prev": has_prev,
        "has_next": has_next,
        "manifest": manifest,
    }


@app.get("/jiandu/gallery/items/{item_id}")
async def get_jiandu_gallery_item_detail(item_id: str, request: Request):
    record = _get_jiandu_gallery_item(item_id)
    return {"item": _serialize_jiandu_gallery_item(record, request, include_boxes=True)}


@app.get("/jiandu/gallery/items/{item_id}/image")
async def get_jiandu_gallery_item_image(item_id: str):
    record = _get_jiandu_gallery_item(item_id)
    image_path = record["_image_path"]
    if not image_path.is_file():
        raise HTTPException(status_code=404, detail=f"Jiandu gallery image not found: {image_path.name}")
    return FileResponse(
        path=image_path,
        media_type="image/bmp",
        filename=record.get("image_file") or image_path.name,
        headers={"Cache-Control": "public, max-age=86400"},
    )


@app.get("/jiandu/gallery/items/{item_id}/thumbnail")
async def get_jiandu_gallery_item_thumbnail(item_id: str, width: int = 480):
    return _render_jiandu_gallery_thumbnail(item_id, width=width)


@app.get("/ocr/jobs")
async def list_jobs(
    q: str = "",
    status: str = "all",
    ocr_model: str = "all",
    created_from: str = "",
    created_to: str = "",
    page: int | None = None,
    page_size: int | None = None,
):
    with jobs_lock:
        mem_jobs = list(jobs.values())
    mem_ids = {job.id for job in mem_jobs}

    archived_rows = _db_query(
        """
        SELECT
            job_id,
            MIN(created_at) AS created_at,
            MAX(COALESCE(updated_at, created_at)) AS updated_at,
            COUNT(*) AS results_count
        FROM documents
        GROUP BY job_id
        ORDER BY created_at DESC
        """
    )

    archived_jobs: list[OCRJob] = []
    archived_counts: dict[str, int] = {}
    for row in archived_rows:
        job_id = row["job_id"]
        if job_id in mem_ids:
            continue
        job = _load_archived_job(job_id, include_results=False)
        if not job:
            continue
        archived_jobs.append(job)
        archived_counts[job_id] = row["results_count"] or 0

    all_jobs = mem_jobs + archived_jobs
    ordered = sorted(all_jobs, key=lambda item: item.created_at, reverse=True)
    job_map = {job.id: job for job in ordered}
    payload_jobs = [
        _job_to_dict(
            job,
            include_results=False,
            results_count=archived_counts.get(job.id),
        )
        for job in ordered
    ]
    excerpts_from_db = _load_job_text_excerpts([payload["id"] for payload in payload_jobs])
    for payload in payload_jobs:
        job_id = payload["id"]
        excerpt = excerpts_from_db.get(job_id, "")
        if not excerpt:
            mem_job = job_map.get(job_id)
            if mem_job:
                excerpt = _truncate_text_excerpt(_extract_preview_text_from_job(mem_job))
        payload["text_excerpt"] = excerpt

    keyword = q.strip().lower()
    status_filter = status.strip().lower() or "all"
    model_filter = ""
    if ocr_model.strip().lower() not in {"", "all"}:
        model_filter = _resolve_ocr_model(ocr_model, strict=False)
    start_dt = _parse_filter_datetime(created_from, end_of_day=False)
    end_dt = _parse_filter_datetime(created_to, end_of_day=True)
    start_ts = start_dt.timestamp() if start_dt else None
    end_ts = end_dt.timestamp() if end_dt else None

    filtered_jobs = [
        payload
        for payload in payload_jobs
        if _job_matches_filters(payload, keyword, status_filter, start_ts, end_ts)
        and (
            not model_filter
            or _resolve_ocr_model(str(payload.get("ocr_model", ""))) == model_filter
        )
    ]

    use_pagination = page is not None or page_size is not None
    if not use_pagination:
        return {
            "jobs": filtered_jobs,
            "total": len(filtered_jobs),
            "page": 1,
            "page_size": len(filtered_jobs) if filtered_jobs else 0,
            "total_pages": 1 if filtered_jobs else 0,
            "has_prev": False,
            "has_next": False,
        }

    safe_page, safe_page_size = _normalize_pagination(
        page,
        page_size,
        default_page_size=20,
        max_page_size=100,
    )
    paged_jobs, total, total_pages, has_prev, has_next = _paginate_items(
        filtered_jobs,
        safe_page,
        safe_page_size,
    )

    current_page = safe_page
    if total_pages == 0:
        current_page = 1
    elif current_page > total_pages:
        current_page = total_pages

    return {
        "jobs": paged_jobs,
        "total": total,
        "page": current_page,
        "page_size": safe_page_size,
        "total_pages": total_pages,
        "has_prev": has_prev,
        "has_next": has_next,
    }


@app.get("/ocr/queue")
async def get_queue():
    return _queue_snapshot()


@app.post("/ocr/queue/reorder")
async def reorder_queue(payload: QueueReorderRequest):
    desired_order = payload.order or []
    with queue_lock:
        existing = list(job_queue)
        new_queue = [job_id for job_id in desired_order if job_id in existing]
        for job_id in existing:
            if job_id not in new_queue:
                new_queue.append(job_id)
        job_queue.clear()
        job_queue.extend(new_queue)
        if job_queue:
            queue_event.set()
        else:
            queue_event.clear()
    return _queue_snapshot()


@app.delete("/ocr/queue/{job_id}")
async def delete_queue_job(job_id: str):
    removed = _remove_from_queue(job_id)
    if not removed:
        raise HTTPException(status_code=404, detail="Job not in queue.")
    with jobs_lock:
        job = jobs.get(job_id)
        if job:
            job.status = "canceled"
            job.pause_requested = False
            job.updated_at = time.time()
            _publish_job_event(
                job_id,
                {"type": "status", "data": {"status": "canceled"}},
            )
            return _job_to_dict(job)
    return {"id": job_id, "status": "canceled"}


@app.post("/ocr/jobs")
async def create_job(
    files: list[UploadFile] = File(...),
    prompt: str = Form(PROMPT_DEFAULT),
    base_size: int = Form(1024),
    image_size: int = Form(768),
    crop_mode: bool = Form(True),
    ocr_model: str = Form(OCR_MODEL_DEEPJIANDU),
    markdown_ignore_labels: list[str] | None = Form(None),
    use_doc_orientation_classify: bool = Form(False),
    use_doc_unwarping: bool = Form(False),
    use_layout_detection: bool = Form(True),
    use_chart_recognition: bool = Form(False),
    layout_threshold: float = Form(0.5),
    layout_unclip_ratio: float = Form(1.0),
    layout_merge_bboxes_mode: str = Form("large"),
    merge_tables: bool = Form(True),
    relevel_titles: bool = Form(True),
    layout_shape_mode: str = Form("auto"),
    prompt_label: str = Form("ocr"),
    repetition_penalty: float = Form(1.0),
    temperature: float = Form(0.0),
    top_p: float = Form(1.0),
    min_pixels: int = Form(147384),
    max_pixels: int = Form(2822400),
    show_formula_number: bool = Form(False),
    prettify_markdown: bool = Form(False),
    visualize: bool = Form(True),
    layout_nms: bool = Form(True),
    restructure_pages: bool = Form(False),
    jiandu_score_thresh: float = Form(0.30),
    jiandu_topk_detect: int = Form(300),
    jiandu_nms_iou: float = Form(0.35),
    jiandu_min_box_size: float = Form(8.0),
    jiandu_max_aspect: float = Form(1.8),
    jiandu_rec_top_k: int = Form(5),
):
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded.")

    job_id = uuid.uuid4().hex
    input_dir = INPUT_ROOT / job_id
    input_dir.mkdir(parents=True, exist_ok=True)

    saved_paths: list[Path] = []
    for upload in files:
        if not upload.filename:
            continue
        dest = input_dir / Path(upload.filename).name
        with dest.open("wb") as buffer:
            shutil.copyfileobj(upload.file, buffer)
        saved_paths.append(dest)

    if not saved_paths:
        raise HTTPException(status_code=400, detail="No valid files saved.")

    resolved_model = _resolve_ocr_model(ocr_model, strict=True)
    normalized_ignore_labels = _normalize_paddle_markdown_ignore_labels(markdown_ignore_labels)
    normalized_layout_shape_mode = _normalize_paddle_layout_shape_mode(layout_shape_mode, "auto")
    normalized_prompt_label = _normalize_paddle_prompt_label(prompt_label, "ocr")
    normalized_layout_merge_bboxes_mode = _normalize_paddle_layout_merge_mode(
        layout_merge_bboxes_mode, "large"
    )
    normalized_layout_threshold = _coerce_float_option(layout_threshold, 0.5, min_value=0.0)
    normalized_layout_threshold = min(1.0, normalized_layout_threshold)
    normalized_layout_unclip_ratio = _coerce_float_option(
        layout_unclip_ratio, 1.0, min_value=0.000001
    )
    normalized_min_pixels = _coerce_int_option(min_pixels, 147384, min_value=1)
    normalized_max_pixels = _coerce_int_option(max_pixels, 2822400, min_value=1)
    if normalized_min_pixels > normalized_max_pixels:
        normalized_min_pixels, normalized_max_pixels = (
            normalized_max_pixels,
            normalized_min_pixels,
        )
    normalized_jiandu_options = normalize_deepjiandu_options(
        {
            "jiandu_score_thresh": jiandu_score_thresh,
            "jiandu_topk_detect": jiandu_topk_detect,
            "jiandu_nms_iou": jiandu_nms_iou,
            "jiandu_min_box_size": jiandu_min_box_size,
            "jiandu_max_aspect": jiandu_max_aspect,
            "jiandu_rec_top_k": jiandu_rec_top_k,
        }
    )

    job_options: dict[str, Any] = {
        "ocr_model": resolved_model,
        "base_size": base_size,
        "image_size": image_size,
        "crop_mode": crop_mode,
    }
    if resolved_model == OCR_MODEL_PADDLE:
        job_options.update(
            {
                "markdown_ignore_labels": normalized_ignore_labels,
                "use_doc_orientation_classify": _coerce_bool_option(
                    use_doc_orientation_classify, False
                ),
                "use_doc_unwarping": _coerce_bool_option(use_doc_unwarping, False),
                "use_layout_detection": _coerce_bool_option(use_layout_detection, True),
                "use_chart_recognition": _coerce_bool_option(use_chart_recognition, False),
                "layout_threshold": normalized_layout_threshold,
                "layout_unclip_ratio": normalized_layout_unclip_ratio,
                "layout_merge_bboxes_mode": normalized_layout_merge_bboxes_mode,
                "merge_tables": _coerce_bool_option(merge_tables, True),
                "relevel_titles": _coerce_bool_option(relevel_titles, True),
                "layout_shape_mode": normalized_layout_shape_mode,
                "prompt_label": normalized_prompt_label,
                "repetition_penalty": _coerce_float_option(
                    repetition_penalty, 1.0, min_value=0.0
                ),
                "temperature": _coerce_float_option(temperature, 0.0, min_value=0.0),
                "top_p": _coerce_float_option(top_p, 1.0, min_value=0.0),
                "min_pixels": normalized_min_pixels,
                "max_pixels": normalized_max_pixels,
                "show_formula_number": _coerce_bool_option(show_formula_number, False),
                "prettify_markdown": _coerce_bool_option(prettify_markdown, False),
                "visualize": _coerce_bool_option(visualize, True),
                "layout_nms": _coerce_bool_option(layout_nms, True),
                "restructure_pages": _coerce_bool_option(restructure_pages, False),
            }
        )
    elif resolved_model == OCR_MODEL_DEEPJIANDU:
        job_options.update(
            {
                "jiandu_score_thresh": normalized_jiandu_options["score_thresh"],
                "jiandu_topk_detect": normalized_jiandu_options["topk_detect"],
                "jiandu_nms_iou": normalized_jiandu_options["nms_iou"],
                "jiandu_min_box_size": normalized_jiandu_options["min_box_size"],
                "jiandu_max_aspect": normalized_jiandu_options["max_aspect"],
                "jiandu_rec_top_k": normalized_jiandu_options["rec_top_k"],
            }
        )
    job = OCRJob(
        id=job_id,
        files=[path.name for path in saved_paths],
        input_paths=[str(path) for path in saved_paths],
        prompt=prompt,
        options=job_options,
        status="queued",
        page_count=len(saved_paths),
    )

    with jobs_lock:
        jobs[job_id] = job

    _enqueue_job(job_id)

    return _job_to_dict(job)


@app.get("/ocr/jobs/{job_id}")
async def get_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
    if job:
        queue_pos = _queue_position(job_id)
        return _job_to_dict(job, queue_position=queue_pos, include_results=False)

    archived_job = _load_archived_job(job_id, include_results=False)
    if not archived_job:
        raise HTTPException(status_code=404, detail="Job not found.")
    return _job_to_dict(archived_job, include_results=False)


@app.get("/ocr/jobs/{job_id}/results")
async def get_job_results(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
    if job:
        results = job.results
        _hydrate_results_from_db(job_id, results)
        return {"results": results, "status": job.status, "error": job.error}

    archived_job = _load_archived_job(job_id, include_results=True)
    if not archived_job:
        raise HTTPException(status_code=404, detail="Job not found.")
    return {"results": archived_job.results, "status": "succeeded", "error": None}


@app.post("/ocr/jobs/{job_id}/pause")
async def pause_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")
        if job.status in {"succeeded", "failed", "canceled"}:
            return _job_to_dict(job)
        if job.status in {"queued", "pending"}:
            job.pause_requested = False
            job.status = "paused"
            job.updated_at = time.time()
            _remove_from_queue(job_id)
            _publish_job_event(
                job_id,
                {"type": "status", "data": {"status": "paused", "progress": job.progress}},
            )
            return _job_to_dict(job)

        job.pause_requested = True
        if job.status == "running":
            job.status = "pausing"
        job.updated_at = time.time()
        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": job.status, "progress": job.progress}},
        )
        return _job_to_dict(job)


@app.post("/ocr/jobs/{job_id}/resume")
async def resume_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")
        if job.status in {"succeeded", "failed", "canceled"}:
            return _job_to_dict(job)
        job.pause_requested = False
        if job.status in {"paused"}:
            job.status = "queued"
            job.updated_at = time.time()
            _enqueue_job(job_id)
            _publish_job_event(
                job_id,
                {"type": "status", "data": {"status": "queued", "progress": job.progress}},
            )
            return _job_to_dict(job)
        if job.status == "pausing":
            job.status = "running"
        job.updated_at = time.time()
        _publish_job_event(
            job_id,
            {"type": "status", "data": {"status": job.status, "progress": job.progress}},
        )
        return _job_to_dict(job)


@app.get("/ocr/jobs/{job_id}/files/{file_path:path}")
async def get_job_file(job_id: str, file_path: str):
    base_dir = (OUTPUT_ROOT / job_id).resolve()
    target = (base_dir / file_path).resolve()
    if base_dir not in target.parents and target != base_dir:
        raise HTTPException(status_code=403, detail="Invalid file path.")
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found.")
    return FileResponse(target)


@app.get("/ocr/jobs/{job_id}/inputs/{file_path:path}")
async def get_job_input(job_id: str, file_path: str):
    base_dir = (INPUT_ROOT / job_id).resolve()
    target = (base_dir / file_path).resolve()
    if base_dir not in target.parents and target != base_dir:
        raise HTTPException(status_code=403, detail="Invalid file path.")
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found.")
    return FileResponse(target)


@app.post("/ocr/jobs/{job_id}/cancel")
async def cancel_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")
        if job.status in {"running", "pausing"}:
            raise HTTPException(status_code=409, detail="Job is running; pause it first.")
        if job.status in {"succeeded", "failed", "canceled"}:
            return _job_to_dict(job)
        job.status = "canceled"
        job.pause_requested = False
        job.updated_at = time.time()
    _remove_from_queue(job_id)
    _publish_job_event(
        job_id,
        {"type": "status", "data": {"status": "canceled"}},
    )
    return _job_to_dict(job)


@app.delete("/ocr/jobs/{job_id}")
async def delete_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if job and job.status in {"running", "pausing"}:
            raise HTTPException(
                status_code=409,
                detail="Job is running; pause it first before deletion.",
            )
        removed_job = jobs.pop(job_id, None)

    _remove_from_queue(job_id)

    doc_rows = _db_query("SELECT id FROM documents WHERE job_id = ?", (job_id,))
    doc_ids = [int(row["id"]) for row in doc_rows]
    deleted_doc_count = len(doc_ids)

    global FTS_ENABLED
    if doc_ids and FTS_ENABLED:
        placeholders = ",".join("?" for _ in doc_ids)
        try:
            _db_execute(
                f"DELETE FROM documents_fts WHERE rowid IN ({placeholders})",
                tuple(doc_ids),
            )
        except sqlite3.Error:
            FTS_ENABLED = False

    if deleted_doc_count > 0:
        _db_execute("DELETE FROM documents WHERE job_id = ?", (job_id,))

    input_dir = INPUT_ROOT / job_id
    output_dir = OUTPUT_ROOT / job_id
    had_input = input_dir.exists()
    had_output = output_dir.exists()
    if had_input:
        shutil.rmtree(input_dir, ignore_errors=True)
    if had_output:
        shutil.rmtree(output_dir, ignore_errors=True)

    with job_streams_lock:
        job_streams.pop(job_id, None)

    deleted_any = bool(removed_job) or deleted_doc_count > 0 or had_input or had_output
    if not deleted_any:
        raise HTTPException(status_code=404, detail="Job not found.")

    return {
        "id": job_id,
        "deleted": True,
        "deleted_documents": deleted_doc_count,
        "deleted_input_dir": had_input,
        "deleted_output_dir": had_output,
    }


@app.get("/ocr/jobs/{job_id}/stream")
async def stream_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")

    q = _register_stream(job_id)

    def event_stream():
        try:
            with jobs_lock:
                job_snapshot = jobs.get(job_id)
            if job_snapshot:
                _hydrate_results_from_db(job_id, job_snapshot.results)
                snapshot = {
                    "type": "snapshot",
                    "data": {
                        "job": _job_to_dict(job_snapshot),
                        "results": job_snapshot.results,
                    },
                }
                yield f"data: {json.dumps(snapshot, ensure_ascii=False)}\n\n"

            while True:
                try:
                    event = q.get(timeout=15)
                except queue.Empty:
                    yield ":\n\n"
                    continue
                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"
        finally:
            _unregister_stream(job_id, q)

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@app.post("/assistant/chat")
async def assistant_chat(request: Request, payload: AssistantChatRequest):
    normalized_messages = _normalize_assistant_messages(payload.messages)
    if not normalized_messages:
        raise HTTPException(status_code=400, detail="秦简问答至少需要一条有效消息。")

    last_user_message = _get_last_user_message(normalized_messages)
    if not last_user_message:
        raise HTTPException(status_code=400, detail="请至少提供一条用户消息。")

    outbound_messages = [{"role": "system", "content": ASSISTANT_SYSTEM_PROMPT}]
    context_payload = {"documents": [], "gallery_items": []}
    if payload.include_platform_context:
        platform_context_text, context_payload = _build_assistant_platform_context(
            last_user_message,
            request,
        )
        outbound_messages.append({"role": "system", "content": platform_context_text})
    outbound_messages.extend(normalized_messages)

    completion = _call_moonshot_chat_completion(
        outbound_messages,
        temperature=payload.temperature,
    )
    return {
        "assistant_name": DEEPBOOK_ASSISTANT_NAME,
        "provider": "moonshot",
        "model": completion["model"],
        "assistant": completion["assistant"],
        "context": context_payload,
        "usage": completion["usage"],
    }


@app.get("/search")
async def search_documents(
    q: str = "",
    limit: int | None = None,
    page: int | None = None,
    page_size: int | None = None,
    job_id: str = "",
    file_name: str = "",
    created_from: str = "",
    created_to: str = "",
):
    query = q.strip()
    normalized_job_filter = job_id.strip().lower()
    normalized_file_filter = file_name.strip().lower()
    start_dt = _parse_filter_datetime(created_from, end_of_day=False)
    end_dt = _parse_filter_datetime(created_to, end_of_day=True)
    start_ts = start_dt.timestamp() if start_dt else None
    end_ts = end_dt.timestamp() if end_dt else None
    derived_page_size = page_size if page_size is not None else limit
    safe_page, safe_page_size = _normalize_pagination(
        page,
        derived_page_size,
        default_page_size=20,
        max_page_size=50,
    )

    if query:
        count_rows = _db_query("SELECT COUNT(*) AS total FROM documents")
        total_docs = int(count_rows[0]["total"]) if count_rows else 0
        fetch_limit = max(1, total_docs)

        try:
            candidates = _search_documents(query, fetch_limit)
        except sqlite3.Error:
            global FTS_ENABLED
            FTS_ENABLED = False
            candidates = _search_documents(query, fetch_limit)

        filtered_candidates = [
            row
            for row in candidates
            if _document_matches_filters(
                row,
                normalized_job_filter,
                normalized_file_filter,
                start_ts,
                end_ts,
            )
        ]

        paged_rows, total, total_pages, has_prev, has_next = _paginate_items(
            filtered_candidates,
            safe_page,
            safe_page_size,
        )

        current_page = safe_page
        if total_pages == 0:
            current_page = 1
        elif current_page > total_pages:
            current_page = total_pages

        return {
            "query": query,
            "filters": {
                "job_id": job_id.strip(),
                "file_name": file_name.strip(),
                "created_from": created_from.strip(),
                "created_to": created_to.strip(),
            },
            "results": [
                {
                    "doc_id": row["id"],
                    "job_id": row["job_id"],
                    "title": row["title"] or row["file_name"] or "",
                    "file_name": row["file_name"],
                    "created_at": row.get("created_at"),
                    "snippet": row.get("snippet") or "",
                    "score": row.get("score", 0.0),
                }
                for row in paged_rows
            ],
            "total": total,
            "page": current_page,
            "page_size": safe_page_size,
            "total_pages": total_pages,
            "has_prev": has_prev,
            "has_next": has_next,
        }

    conditions: list[str] = []
    params: list[Any] = []
    if normalized_job_filter:
        conditions.append("job_id LIKE ?")
        params.append(f"%{job_id.strip()}%")
    if normalized_file_filter:
        conditions.append("(file_name LIKE ? OR title LIKE ?)")
        like = f"%{file_name.strip()}%"
        params.extend([like, like])
    if start_dt:
        conditions.append("created_at >= ?")
        params.append(start_dt.isoformat())
    if end_dt:
        conditions.append("created_at <= ?")
        params.append(end_dt.isoformat())

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    count_query = f"SELECT COUNT(*) AS total FROM documents {where_clause}"
    total_rows = _db_query(count_query, tuple(params))
    total = int(total_rows[0]["total"]) if total_rows else 0
    total_pages = (total + safe_page_size - 1) // safe_page_size if total > 0 else 0

    current_page = safe_page
    if total_pages == 0:
        current_page = 1
    elif current_page > total_pages:
        current_page = total_pages
    offset = (current_page - 1) * safe_page_size
    has_prev = total_pages > 0 and current_page > 1
    has_next = total_pages > 0 and current_page < total_pages

    rows = _db_query(
        f"""
        SELECT
            id, job_id, title, file_name, created_at,
            markdown, text, raw_text, edited_markdown, edited_text
        FROM documents
        {where_clause}
        ORDER BY created_at DESC, id DESC
        LIMIT ? OFFSET ?
        """,
        tuple([*params, safe_page_size, offset]),
    )
    results = []
    for row in rows:
        content = _coalesce_document_content(row)
        results.append(
            {
                "doc_id": row["id"],
                "job_id": row["job_id"],
                "title": row["title"] or row["file_name"] or "",
                "file_name": row["file_name"],
                "created_at": row["created_at"],
                "snippet": _build_plain_preview(content),
                "score": 0.0,
            }
        )
    return {
        "query": query,
        "filters": {
            "job_id": job_id.strip(),
            "file_name": file_name.strip(),
            "created_from": created_from.strip(),
            "created_to": created_to.strip(),
        },
        "results": results,
        "total": total,
        "page": current_page,
        "page_size": safe_page_size,
        "total_pages": total_pages,
        "has_prev": has_prev,
        "has_next": has_next,
    }


@app.get("/documents/{doc_id}")
async def get_document(doc_id: int):
    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text,
            output_dir, output_dir_rel, output_files, boxes_image,
            cropped_images
        FROM documents
        WHERE id = ?
        """,
        (doc_id,),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Document not found.")
    return {"document": _row_to_document(rows[0])}


@app.put("/documents/{doc_id}")
async def update_document(doc_id: int, payload: DocumentUpdateRequest):
    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text
        FROM documents
        WHERE id = ?
        """,
        (doc_id,),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Document not found.")
    row = rows[0]
    provided_markdown = payload.edited_markdown is not None
    provided_text = payload.edited_text is not None
    edited_markdown = (
        payload.edited_markdown if provided_markdown else row["edited_markdown"]
    )
    edited_text = payload.edited_text if provided_text else row["edited_text"]

    if not provided_markdown and not provided_text:
        raise HTTPException(status_code=400, detail="No content provided.")

    if provided_markdown and not provided_text:
        edited_text = _markdown_to_text(payload.edited_markdown or "")

    updated_at = datetime.now(tz=timezone.utc).isoformat()
    _db_execute(
        """
        UPDATE documents
        SET edited_markdown = ?, edited_text = ?, updated_at = ?
        WHERE id = ?
        """,
        (edited_markdown, edited_text, updated_at, doc_id),
    )

    title = row["title"] or row["file_name"] or ""
    if edited_text is not None:
        content = edited_text
    elif edited_markdown is not None:
        content = edited_markdown
    else:
        content = row["text"] or row["raw_text"] or ""
    _upsert_document_fts(doc_id, title, content)

    with jobs_lock:
        job = jobs.get(row["job_id"])
        if job:
            for result in job.results:
                if result.get("doc_id") == doc_id or (
                    row["file_name"] and result.get("file") == row["file_name"]
                ):
                    result["edited_markdown"] = edited_markdown
                    result["edited_text"] = edited_text
                    break

    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text,
            output_dir, output_dir_rel, output_files, boxes_image,
            cropped_images
        FROM documents
        WHERE id = ?
        """,
        (doc_id,),
    )
    return {"document": _row_to_document(rows[0])}


def _resolve_output_dir_rel_for_row(row: sqlite3.Row) -> str:
    output_dir_rel = row["output_dir_rel"] or ""
    output_dir = row["output_dir"] or ""
    job_id = row["job_id"] or ""
    if output_dir_rel or not output_dir or not job_id:
        return output_dir_rel
    try:
        return _relative_to_job(job_id, Path(output_dir)) or ""
    except Exception:
        return ""


def _parse_export_doc_ids(raw_doc_ids: str) -> list[int]:
    text = (raw_doc_ids or "").strip()
    if not text:
        return []
    parsed: list[int] = []
    seen: set[int] = set()
    for chunk in text.split(","):
        token = chunk.strip()
        if not token:
            continue
        if not token.isdigit():
            raise HTTPException(status_code=400, detail="doc_ids must be comma-separated integers.")
        doc_id = int(token)
        if doc_id <= 0:
            raise HTTPException(status_code=400, detail="doc_ids must be positive integers.")
        if doc_id in seen:
            continue
        seen.add(doc_id)
        parsed.append(doc_id)
    return parsed


def _prefix_markdown_asset_url(url: str, output_dir_rel: str) -> str:
    if not url:
        return url
    raw = url.strip()
    wrapped = raw.startswith("<") and raw.endswith(">")
    if wrapped:
        raw = raw[1:-1].strip()
    if not raw or raw.startswith("#"):
        return url
    parsed = urlsplit(raw)
    if parsed.scheme or parsed.netloc:
        return url

    rel_path = _normalize_markdown_asset_subpath(parsed.path)
    if not rel_path:
        return url

    base_rel = (output_dir_rel or "").replace("\\", "/").strip("/")
    if base_rel:
        if rel_path.startswith(f"{base_rel}/"):
            suffix = _normalize_markdown_asset_subpath(rel_path[len(base_rel) + 1 :])
            rel_path = f"{base_rel}/{suffix}" if suffix else base_rel
        else:
            rel_path = f"{base_rel}/{rel_path}"

    rebuilt = rel_path
    if parsed.query:
        rebuilt = f"{rebuilt}?{parsed.query}"
    if parsed.fragment:
        rebuilt = f"{rebuilt}#{parsed.fragment}"
    if wrapped:
        return f"<{rebuilt}>"
    return rebuilt


def _prefix_markdown_asset_links(markdown: str, output_dir_rel: str) -> str:
    if not markdown:
        return markdown

    def replace_markdown_image(match: re.Match[str]) -> str:
        alt_text = match.group(1)
        inner = match.group(2).strip()
        if not inner:
            return match.group(0)

        link_part = inner
        suffix = ""
        if inner.startswith("<"):
            closing_idx = inner.find(">")
            if closing_idx != -1:
                link_part = inner[: closing_idx + 1]
                suffix = inner[closing_idx + 1 :].strip()
        else:
            parts = inner.split(maxsplit=1)
            link_part = parts[0]
            suffix = parts[1] if len(parts) > 1 else ""

        rewritten = _prefix_markdown_asset_url(link_part, output_dir_rel)
        if rewritten == link_part:
            return match.group(0)
        rebuilt = rewritten + (f" {suffix}" if suffix else "")
        return f"![{alt_text}]({rebuilt})"

    output = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", replace_markdown_image, markdown)

    def replace_html_image(match: re.Match[str]) -> str:
        prefix, link, suffix = match.groups()
        rewritten = _prefix_markdown_asset_url(link, output_dir_rel)
        return f"{prefix}{rewritten}{suffix}"

    return re.sub(
        r'(<img\b[^>]*\bsrc=["\'])([^"\']+)(["\'])',
        replace_html_image,
        output,
        flags=re.IGNORECASE,
    )


def _export_bundle(
    fmt: str,
    title: str,
    markdown: str,
    text: str,
    resource_paths: list[str] | None = None,
) -> Response:
    normalized_fmt = fmt.lower()
    markdown_content = markdown or text or ""
    text_content = text or markdown_content or ""
    normalized_math_markdown_content = _normalize_export_math_delimiters(markdown_content)
    normalized_markdown_content = _normalize_markdown_asset_links(normalized_math_markdown_content)
    centered_markdown_content = _center_markdown_figure_captions(normalized_markdown_content)
    rich_markdown_content = _convert_html_images_to_markdown(centered_markdown_content)
    rich_markdown_content = _normalize_export_math_delimiters(rich_markdown_content)

    if normalized_fmt in {"md", "markdown"}:
        styled_markdown_content = _style_standalone_images_for_markdown(
            centered_markdown_content, default_width="88%"
        )
        markdown_for_export = _normalize_markdown_math_for_md(styled_markdown_content)
        filename = f"{title}.md"
        return Response(
            content=markdown_for_export,
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": _content_disposition(filename)},
        )

    if normalized_fmt in {"txt", "text"}:
        filename = f"{title}.txt"
        return Response(
            content=text_content,
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": _content_disposition(filename)},
        )

    if normalized_fmt == "docx":
        pandoc_bytes, pandoc_error = _export_with_pandoc(
            rich_markdown_content, "docx", resource_paths
        )
        if pandoc_bytes:
            alt_candidates = _collect_markdown_image_alts(rich_markdown_content)
            docx_bytes = _center_docx_image_alt_paragraphs(pandoc_bytes, alt_candidates)
            filename = f"{title}.docx"
            return Response(
                content=docx_bytes,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": _content_disposition(filename)},
            )
        if _markdown_has_images(rich_markdown_content):
            detail = (
                "Pandoc is required to export images in DOCX. "
                "Install pandoc and ensure it is available in PATH."
            )
            extra = _short_error_message(pandoc_error)
            if extra:
                detail = f"{detail} Conversion error: {extra}"
            raise HTTPException(
                status_code=500,
                detail=detail,
            )
        try:
            from docx import Document as DocxDocument
        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail="python-docx not installed. Install python-docx to export docx.",
            ) from exc

        doc = DocxDocument()
        for line in text_content.splitlines():
            doc.add_paragraph(line)
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        filename = f"{title}.docx"
        return Response(
            content=buffer.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": _content_disposition(filename)},
        )

    if normalized_fmt == "pdf":
        pdf_markdown_content = _remove_markdown_image_captions(rich_markdown_content)
        pdf_markdown_for_pandoc = _convert_standalone_markdown_images_to_latex(
            pdf_markdown_content
        )
        pandoc_bytes, pandoc_error = _export_with_pandoc(
            pdf_markdown_for_pandoc, "pdf", resource_paths
        )
        if pandoc_bytes:
            return _pdf_response(title, pandoc_bytes)
        return _export_pdf_with_reportlab(
            title=title,
            markdown_content=pdf_markdown_content,
            text_content=text_content,
            resource_paths=resource_paths,
            pandoc_error=pandoc_error,
        )

    raise HTTPException(status_code=400, detail="Unsupported export format.")


@app.get("/documents/{doc_id}/export")
async def export_document(
    doc_id: int,
    request: Request,
    format: str = "md",
    asset_base: str | None = None,
):
    rows = _db_query(
        """
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text, output_dir, output_dir_rel
        FROM documents
        WHERE id = ?
        """,
        (doc_id,),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Document not found.")

    row = rows[0]
    title = row["title"] or row["file_name"] or f"document_{doc_id}"
    markdown = (
        row["edited_markdown"] if row["edited_markdown"] is not None else row["markdown"] or ""
    )
    text = (
        row["edited_text"]
        if row["edited_text"] is not None
        else row["text"] or row["raw_text"] or markdown or ""
    )
    output_dir = None
    try:
        output_dir = row["output_dir"]
    except Exception:
        output_dir = None
    job_id = row["job_id"] or ""
    output_dir_rel = row["output_dir_rel"] or ""
    if not output_dir_rel and output_dir and job_id:
        try:
            output_dir_rel = _relative_to_job(job_id, Path(output_dir)) or ""
        except Exception:
            output_dir_rel = ""
    resource_paths: list[str] = []
    if output_dir:
        resource_paths.append(output_dir)
        images_dir = str(Path(output_dir) / "images")
        resource_paths.append(images_dir)

    fmt = format.lower()
    export_markdown = markdown or text
    if fmt in {"md", "markdown"} and export_markdown and job_id:
        public_api_base = _resolve_public_api_base(request, asset_base)
        export_markdown = _rewrite_markdown_asset_links(
            export_markdown, job_id, output_dir_rel, public_api_base
        )
    return _export_bundle(fmt, title, export_markdown, text, resource_paths)


@app.get("/ocr/jobs/{job_id}/export")
async def export_job_documents(
    job_id: str,
    request: Request,
    format: str = "md",
    doc_ids: str = "",
    asset_base: str | None = None,
):
    selected_doc_ids = _parse_export_doc_ids(doc_ids)
    fmt = format.lower()

    params: list[Any] = [job_id]
    where_clause = "job_id = ?"
    if selected_doc_ids:
        placeholders = ",".join("?" for _ in selected_doc_ids)
        where_clause = f"{where_clause} AND id IN ({placeholders})"
        params.extend(selected_doc_ids)

    rows = _db_query(
        f"""
        SELECT
            id, job_id, title, file_name, markdown, text, raw_text,
            edited_markdown, edited_text, output_dir, output_dir_rel
        FROM documents
        WHERE {where_clause}
        ORDER BY id
        """,
        tuple(params),
    )

    if not rows:
        raise HTTPException(status_code=404, detail="No exportable pages found for this job.")

    if selected_doc_ids:
        found_ids = {int(row["id"]) for row in rows}
        missing_ids = [doc_id for doc_id in selected_doc_ids if doc_id not in found_ids]
        if missing_ids:
            raise HTTPException(
                status_code=404,
                detail=f"Some selected pages are not found in this job: {missing_ids}",
            )
        order_map = {doc_id: index for index, doc_id in enumerate(selected_doc_ids)}
        rows = sorted(rows, key=lambda row: order_map.get(int(row["id"]), len(order_map)))

    public_api_base = ""
    if fmt in {"md", "markdown"}:
        public_api_base = _resolve_public_api_base(request, asset_base)

    resource_paths: set[str] = set()
    job_output_dir = OUTPUT_ROOT / job_id
    if job_output_dir.exists():
        resource_paths.add(str(job_output_dir))

    markdown_sections: list[str] = []
    markdown_local_sections: list[str] = []
    text_sections: list[str] = []
    per_page_pdf_entries: list[dict[str, Any]] = []

    for index, row in enumerate(rows, start=1):
        page_title = row["title"] or row["file_name"] or f"第{index}页"
        output_dir = row["output_dir"] or ""
        output_dir_rel = _resolve_output_dir_rel_for_row(row)

        page_markdown = (
            row["edited_markdown"] if row["edited_markdown"] is not None else row["markdown"] or ""
        )
        page_text = (
            row["edited_text"]
            if row["edited_text"] is not None
            else row["text"] or row["raw_text"] or page_markdown or ""
        )
        page_content = page_markdown or page_text
        heading = f"## 第{index}页 · {page_title}"

        if page_content and public_api_base:
            export_markdown = _rewrite_markdown_asset_links(
                page_content,
                job_id,
                output_dir_rel,
                public_api_base,
            )
        else:
            export_markdown = page_content
        markdown_sections.append(
            f"{heading}\n\n{export_markdown}".strip() if page_content else heading
        )

        local_markdown = _prefix_markdown_asset_links(page_content, output_dir_rel)
        markdown_local_sections.append(
            f"{heading}\n\n{local_markdown}".strip() if page_content else heading
        )
        text_sections.append(f"===== 第{index}页 · {page_title} =====\n{page_text}".strip())

        page_resource_paths: set[str] = set()
        if output_dir:
            page_resource_paths.add(output_dir)
            page_resource_paths.add(str(Path(output_dir) / "images"))
        if output_dir:
            resource_paths.add(output_dir)
            resource_paths.add(str(Path(output_dir) / "images"))
        page_markdown_for_pdf = page_markdown or page_text or ""
        page_text_for_pdf = page_text or page_markdown or ""
        per_page_pdf_entries.append(
            {
                "title": page_title,
                "markdown": page_markdown_for_pdf,
                "text": page_text_for_pdf,
                "resource_paths": sorted(path for path in page_resource_paths if path),
            }
        )

    combined_markdown = "\n\n---\n\n".join(section for section in markdown_sections if section)
    combined_local_markdown = "\n\n---\n\n".join(
        section for section in markdown_local_sections if section
    )
    combined_text = "\n\n".join(section for section in text_sections if section)

    selection_label = "选定页" if selected_doc_ids else "全部页"
    title = f"任务_{job_id[:8]}_{selection_label}_{len(rows)}页"
    resource_path_list = sorted(path for path in resource_paths if path)

    if fmt in {"md", "markdown"}:
        return _export_bundle(fmt, title, combined_markdown, combined_text, resource_path_list)
    if fmt == "pdf":
        page_pdf_chunks: list[bytes] = []
        for page in per_page_pdf_entries:
            page_pdf_response = _export_bundle(
                "pdf",
                page["title"] or "page",
                page["markdown"],
                page["text"],
                page["resource_paths"],
            )
            page_pdf_chunks.append(bytes(page_pdf_response.body or b""))
        if len(page_pdf_chunks) == 1:
            return _pdf_response(title, page_pdf_chunks[0])

        merged_pdf = _merge_pdf_chunks(page_pdf_chunks)
        if merged_pdf:
            return _pdf_response(title, merged_pdf)
        return _export_bundle("pdf", title, combined_local_markdown, combined_text, resource_path_list)
    return _export_bundle(
        fmt,
        title,
        combined_local_markdown,
        combined_text,
        resource_path_list,
    )


if __name__ == "__main__":
    import uvicorn

    multiprocessing.freeze_support()

    host = os.getenv("DEEPBOOK_HOST", "127.0.0.1").strip() or "127.0.0.1"
    port = _coerce_int_option(os.getenv("DEEPBOOK_PORT"), 8000, min_value=1)
    reload_enabled = _coerce_bool_option(os.getenv("DEEPBOOK_RELOAD"), False)

    if reload_enabled:
        uvicorn.run("main:app", host=host, port=port, reload=True)
    else:
        uvicorn.run(app, host=host, port=port)
